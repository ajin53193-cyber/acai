"""Reports & exports domain: analytics xlsx downloads, background export jobs
(GridFS), scheduled exports and the export cron.

Heavy generators + job pipeline extracted from server.py. Data helpers that are
shared with live endpoints (merchant_stats, build_profit_snapshot,
resolve_ledger_account) and label constants remain in server.py and are imported
late (router included at the bottom of server.py).
"""
import asyncio
import io
import os
import re
import hmac
import secrets
import logging
from datetime import datetime, timezone, timedelta
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, Request, Query, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from urllib.parse import quote

from core import db, fs
from utils import serialize, iso, now_utc
from deps import get_current_user, require_roles, merchant_ids_for_scope
from shared import audit, send_telegram, notify_settings
from server import (
    merchant_stats, build_profit_snapshot, resolve_ledger_account,
    ORDER_STATUS_LABELS, LEDGER_CAT_LABELS, LEDGER_CAT, EXPORT_KINDS,
)

logger = logging.getLogger("sifang")
router = APIRouter(prefix="/api")


@router.get("/reports/merchant-stats")
async def merchant_stats_report(date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                                user: dict = Depends(require_roles("admin"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    data = await merchant_stats(date=date, start=start, end=end, user=user)
    wb = Workbook(); ws = wb.active; ws.title = "商户统计"
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws.append([f"统计区间：{data['date']}"])
    ws.append(["商户", "邮箱", "订单数", "成功", "成功率%", "收款金额", "手续费", "实收", "可用余额"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for m in data["merchants"]:
        ws.append([m["name"], m["email"], m["orders"], m["paid"], m["success_rate"],
                   m["volume"], m["fee"], m["receive"], m["balance"]])
    for i, w in enumerate([18, 26, 10, 8, 10, 16, 14, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"merchant_stats_{data['date'].replace(' ', '').replace('~', '-')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@router.get("/reports/agent-commission")
async def agent_commission_report(month: str, agent_id: Optional[str] = None,
                                  user: dict = Depends(require_roles("admin", "agent"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if user["role"] == "agent":
        agent = user
    else:
        if not agent_id:
            raise HTTPException(status_code=400, detail="请选择代理")
        agent = await db.users.find_one({"_id": ObjectId(agent_id), "role": "agent"})
        if not agent:
            raise HTTPException(status_code=404, detail="代理不存在")
    try:
        start = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
    end = (start.replace(year=start.year + 1, month=1) if start.month == 12 else start.replace(month=start.month + 1))

    merchants = await db.users.find({"role": "merchant", "parent_agent_id": str(agent["_id"])}).to_list(500)
    m_ids = [str(m["_id"]) for m in merchants]
    orders = await db.orders.find({"merchant_id": {"$in": m_ids}, "status": "paid",
                                   "paid_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("paid_at", 1).to_list(100000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F766E")

    def write_sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    total = {"n": len(orders), "amt": sum(o["amount_usdt"] for o in orders), "fee": sum(o.get("fee", 0) for o in orders),
             "comm": sum(o.get("agent_commission", 0) for o in orders), "recv": sum(o.get("merchant_receive", 0) for o in orders)}
    ws = wb.active
    ws.title = "月度汇总"
    write_sheet(ws, ["项目", "数值"], [
        ["代理", agent["name"]], ["代理邮箱", agent["email"]], ["月份", month],
        ["代理费率", agent.get("rate", 0)], ["下级商户数", len(merchants)],
        ["成功订单数", total["n"]], ["交易流水 (USDT)", round(total["amt"], 2)],
        ["手续费合计 (USDT)", round(total["fee"], 2)], ["商户实收合计 (USDT)", round(total["recv"], 2)],
        ["代理分润合计 (USDT)", round(total["comm"], 2)],
        ["生成时间", now_utc().strftime("%Y-%m-%d %H:%M UTC")]], [24, 36])

    by_m: dict = {}
    for o in orders:
        d = by_m.setdefault(o["merchant_id"], {"name": o["merchant_name"], "n": 0, "amt": 0.0, "fee": 0.0, "comm": 0.0, "recv": 0.0})
        d["n"] += 1; d["amt"] += o["amount_usdt"]; d["fee"] += o.get("fee", 0); d["comm"] += o.get("agent_commission", 0); d["recv"] += o.get("merchant_receive", 0)
    m_rate = {str(m["_id"]): m.get("rate", 0) for m in merchants}
    write_sheet(wb.create_sheet("按商户明细"), ["商户", "商户费率", "订单数", "交易流水", "手续费", "商户实收", "代理分润"],
                [[d["name"], m_rate.get(mid, ""), d["n"], round(d["amt"], 2), round(d["fee"], 2), round(d["recv"], 2), round(d["comm"], 2)]
                 for mid, d in sorted(by_m.items(), key=lambda x: -x[1]["amt"])], [22, 12, 10, 16, 16, 16, 16])

    write_sheet(wb.create_sheet("逐单明细"), ["订单号", "商户", "网络", "金额", "手续费", "商户实收", "代理分润", "TxID", "支付时间"],
                [[o["order_no"], o["merchant_name"], o["network"], o["amount_usdt"], o.get("fee", 0), o.get("merchant_receive", 0),
                  o.get("agent_commission", 0), o.get("tx_id") or "", (o.get("paid_at") or "")[:19].replace("T", " ")] for o in orders],
                [26, 18, 10, 12, 12, 12, 12, 40, 20])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"commission_{agent['name']}_{month}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@router.get("/reports/merchant-statement")
async def merchant_statement(month: str, merchant_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if user["role"] == "merchant":
        merchant = user
    else:
        if user["role"] == "operator" and "reports.export" not in (user.get("permissions") or []):
            raise HTTPException(status_code=403, detail="无权限访问")
        if not merchant_id:
            raise HTTPException(status_code=400, detail="请选择商户")
        merchant = await db.users.find_one({"_id": ObjectId(merchant_id), "role": "merchant"})
        if not merchant:
            raise HTTPException(status_code=404, detail="商户不存在")
        if user["role"] == "agent" and merchant.get("parent_agent_id") != str(user["_id"]):
            raise HTTPException(status_code=403, detail="无权限")
    try:
        start = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
    end = (start.replace(year=start.year + 1, month=1) if start.month == 12 else start.replace(month=start.month + 1))
    mid = str(merchant["_id"])
    orders = await db.orders.find({"merchant_id": mid, "status": "paid", "paid_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("paid_at", 1).to_list(100000)
    payouts = await db.payouts.find({"merchant_id": mid, "created_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("created_at", 1).to_list(10000)
    adjustments = await db.audit_logs.find({"action": {"$regex": "^balance\\."}, "target_id": mid,
                                            "created_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("created_at", 1).to_list(10000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0E7490")

    def write_sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    amt = sum(o["amount_usdt"] for o in orders); fee = sum(o.get("fee", 0) for o in orders); recv = sum(o.get("merchant_receive", 0) for o in orders)
    approved = sum(p["amount"] for p in payouts if p["status"] == "approved")
    by_net: dict = {}
    for o in orders:
        d = by_net.setdefault(o["network"], {"n": 0, "amt": 0.0, "fee": 0.0})
        d["n"] += 1; d["amt"] += o["amount_usdt"]; d["fee"] += o.get("fee", 0)
    ws = wb.active; ws.title = "月度汇总"
    write_sheet(ws, ["项目", "数值"], [
        ["商户", merchant["name"]], ["商户邮箱", merchant["email"]], ["月份", month], ["商户费率", merchant.get("rate", 0)],
        ["成功订单数", len(orders)], ["收款流水 (USDT)", round(amt, 2)], ["手续费合计 (USDT)", round(fee, 2)], ["实收合计 (USDT)", round(recv, 2)],
        ["提现申请笔数", len(payouts)], ["已通过提现 (USDT)", round(approved, 2)],
        ["当前可用余额 (USDT)", round(merchant.get("balance", 0), 2)], ["当前冻结余额 (USDT)", round(merchant.get("frozen_balance", 0), 2)],
        ["生成时间", now_utc().strftime("%Y-%m-%d %H:%M UTC")]] + [[f"{n} 订单数/流水/手续费", f"{d['n']} / {round(d['amt'], 2)} / {round(d['fee'], 2)}"] for n, d in by_net.items()], [26, 40])
    write_sheet(wb.create_sheet("收款明细"), ["订单号", "商户单号", "网络", "订单金额", "实付金额", "手续费", "实收", "TxID", "收款地址", "支付时间"],
                [[o["order_no"], o.get("out_trade_no") or "", o["network"], o["amount_usdt"], o.get("pay_amount", o["amount_usdt"]), o.get("fee", 0), o.get("merchant_receive", 0),
                  o.get("tx_id") or "", o["wallet_address"], (o.get("paid_at") or "")[:19].replace("T", " ")] for o in orders], [26, 18, 8, 12, 12, 10, 12, 40, 36, 20])
    write_sheet(wb.create_sheet("提现记录"), ["申请时间", "金额", "网络", "提现地址", "状态", "审核时间"],
                [[(p.get("created_at") or "")[:19].replace("T", " "), p["amount"], p.get("network", ""), p.get("address", ""), p["status"], (p.get("approved_at") or "")[:19].replace("T", " ")] for p in payouts], [20, 12, 8, 40, 10, 20])
    write_sheet(wb.create_sheet("余额调整"), ["时间", "操作", "金额", "余额前", "余额后", "操作人", "备注"],
                [[(a["created_at"] or "")[:19].replace("T", " "), a["action"], a["detail"].get("amount"), a["detail"].get("balance_before"), a["detail"].get("balance_after"), a["actor_email"], a["detail"].get("remark", "")] for a in adjustments], [20, 16, 12, 12, 12, 24, 24])

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"statement_{merchant['name']}_{month}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@router.get("/reports/balance-ledger")
async def balance_ledger_report(account_id: Optional[str] = None, month: Optional[str] = None,
                                start: Optional[str] = None, end: Optional[str] = None,
                                user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    acc_id = await resolve_ledger_account(user, account_id)
    acc = await db.users.find_one({"_id": acc_id})
    q: dict = {"account_id": str(acc_id)}
    title_suffix = "全部"
    if start and end:
        try:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
        title_suffix = f"{start}_{end}"
    elif month:
        try:
            ms = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
        me = (ms.replace(year=ms.year + 1, month=1) if ms.month == 12 else ms.replace(month=ms.month + 1))
        q["created_at"] = {"$gte": iso(ms), "$lt": iso(me)}
        title_suffix = month
    rows = await db.balance_ledger.find(q).sort("created_at", 1).to_list(100000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws = wb.active; ws.title = "资金流水"
    ws.append(["时间", "类型", "变动金额(USDT)", "变动后余额", "变动后冻结", "关联", "备注", "操作人"])
    for c in ws[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for r in rows:
        ws.append([(r.get("created_at") or "")[:19].replace("T", " "), LEDGER_CAT_LABELS.get(r["category"], r["category"]),
                   r.get("amount"), r.get("balance_after"), r.get("frozen_after"),
                   f"{r.get('ref_type') or ''} {(r.get('ref_id') or '')}".strip(), r.get("remark", ""), r.get("operator") or ""])
    for i, w in enumerate([20, 14, 16, 14, 14, 30, 34, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"ledger_{acc.get('name', '')}_{title_suffix}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@router.get("/reports/profit-snapshot")
async def profit_snapshot_report(date: str, user: dict = Depends(require_roles("admin"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    snap = await db.profit_snapshots.find_one({"date": date}) or await build_profit_snapshot(date)
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws = wb.active; ws.title = "日度汇总"
    ws.append([f"分润对账日期：{date}"])
    ws.append(["订单数", "收款总额", "手续费", "商户实收", "代理分润", "平台利润", "补偿结算"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    ws.append([snap["order_count"], snap["total_amount"], snap["total_fee"], snap["total_merchant_receive"],
               snap["total_agent_commission"], snap["total_platform_profit"], snap.get("backfilled", 0)])
    for i, w in enumerate([12, 16, 14, 16, 14, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws2 = wb.create_sheet("按商户明细")
    ws2.append(["商户", "订单数", "收款金额", "手续费", "商户实收", "代理分润", "平台利润"])
    for c in ws2[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for m in snap.get("merchants", []):
        ws2.append([m["merchant_name"], m["orders"], m["amount"], m["fee"], m["receive"], m["commission"], m["platform"]])
    for i, w in enumerate([22, 10, 16, 14, 16, 14, 14], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"; ws2.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"profit_snapshot_{date}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


async def generate_orders_xlsx(user, status=None, network=None, keyword=None, merchant_id=None,
                               date=None, start=None, end=None, ids=None, progress=None):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    id_list = [x for x in (ids.split(",") if ids else []) if x]
    if id_list:
        try:
            q["_id"] = {"$in": [ObjectId(x) for x in id_list]}
        except Exception:
            raise HTTPException(status_code=400, detail="订单 ID 格式不正确")
    else:
        if merchant_id and (scope is None or ObjectId(merchant_id) in scope):
            q["merchant_id"] = str(merchant_id)
        if start and end:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
            q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
        elif date:
            ds = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            q["created_at"] = {"$gte": iso(ds), "$lt": iso(ds + timedelta(days=1))}
        if status:
            q["status"] = status
        if network:
            q["network"] = network
        if keyword:
            kw = keyword.strip()
            q["$or"] = [{"order_no": {"$regex": kw, "$options": "i"}}, {"tx_id": {"$regex": kw, "$options": "i"}},
                        {"out_trade_no": {"$regex": kw, "$options": "i"}}, {"wallet_address": {"$regex": kw, "$options": "i"}},
                        {"merchant_name": {"$regex": kw, "$options": "i"}}]
            if re.fullmatch(r"\d+(\.\d+)?", kw):
                q["$or"] += [{"amount_usdt": float(kw)}, {"pay_amount": float(kw)}]
    filters = []
    if id_list:
        filters.append(f"所选订单 {len(id_list)} 笔")
    else:
        if merchant_id:
            filters.append(f"商户ID={merchant_id}")
        if start and end:
            filters.append(f"区间 {start}~{end}")
        elif date:
            filters.append(f"日期 {date}")
        if status:
            filters.append(f"状态={ORDER_STATUS_LABELS.get(status, status)}")
        if network:
            filters.append(f"网络={network}")
        if keyword:
            filters.append(f"关键词={keyword}")
    summary = "、".join(filters) or "全部订单"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("订单明细")
    for i, w in enumerate([22, 16, 18, 8, 12, 12, 10, 36, 40, 10, 12, 10, 20, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    title = WriteOnlyCell(ws, value=f"订单导出 · 筛选条件：{summary}"); title.font = Font(bold=True, size=12)
    ws.append([title])
    ws.append([f"导出时间：{iso(now_utc())[:19].replace('T', ' ')} (UTC)    操作人：{user.get('email', '')}"])
    ws.append([])
    header = ["订单号", "商户", "商户单号", "网络", "金额(USDT)", "到账金额", "状态", "收款地址", "TxID",
              "手续费", "实收", "回调状态", "创建时间", "支付时间"]
    hrow = []
    for h in header:
        c = WriteOnlyCell(ws, value=h); c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        hrow.append(c)
    ws.append(hrow)
    total = await db.orders.count_documents(q)
    if progress:
        await progress(0, total)
    count = 0
    async for o in db.orders.find(q).sort("created_at", -1):
        ws.append([o.get("order_no"), o.get("merchant_name"), o.get("out_trade_no") or "", o.get("network"),
                   o.get("amount_usdt"), o.get("pay_amount"), ORDER_STATUS_LABELS.get(o.get("status"), o.get("status")),
                   o.get("wallet_address") or "", o.get("tx_id") or "", o.get("fee"), o.get("merchant_receive"),
                   o.get("callback_status") or "", (o.get("created_at") or "")[:19].replace("T", " "),
                   (o.get("paid_at") or "")[:19].replace("T", " ")])
        count += 1
        if progress and count % 1000 == 0:
            await progress(count, total)
    ws.freeze_panes = "A5"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    tag = "selected" if id_list else (date or (f"{start}_{end}" if start and end else "all"))
    return buf, count, f"orders_{tag}.xlsx"


@router.get("/reports/orders")
async def orders_report(user: dict = Depends(get_current_user),
                        status: Optional[str] = None, network: Optional[str] = None,
                        keyword: Optional[str] = None, merchant_id: Optional[str] = None,
                        date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                        ids: Optional[str] = None):
    buf, _count, fname = await generate_orders_xlsx(user, status, network, keyword, merchant_id, date, start, end, ids)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


async def generate_snapshots_xlsx(start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    q = {}
    if start and end:
        q["date"] = {"$gte": start, "$lte": end}
    rows = await db.profit_snapshots.find(q).sort("date", -1).to_list(500)
    wb = Workbook(); ws = wb.active; ws.title = "分润快照"
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="0E7490")
    ws.append(["日期", "订单数", "收款总额", "手续费", "商户实收", "代理分润", "平台利润", "补偿单", "环比%"])
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, hfill, Alignment(horizontal="center")
    for s in rows:
        ws.append([s.get("date"), s.get("order_count"), s.get("total_amount"), s.get("total_fee"),
                   s.get("total_merchant_receive"), s.get("total_agent_commission"), s.get("total_platform_profit"),
                   s.get("backfilled", 0), s.get("platform_profit_change_pct")])
    for i, w in enumerate([12, 10, 14, 12, 14, 12, 12, 8, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, len(rows), f"profit_snapshots_{start or 'all'}.xlsx"


async def generate_ledger_xlsx(user, account_id=None, month=None, start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    q = {"account_id": account_id}
    if start and end:
        ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    elif month:
        ds = datetime.strptime(month + "-01", "%Y-%m-%d").replace(tzinfo=timezone.utc)
        de = (ds + timedelta(days=32)).replace(day=1)
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    rows = await db.balance_logs.find(q).sort("created_at", -1).to_list(100000)
    wb = Workbook(); ws = wb.active; ws.title = "资金流水"
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="0E7490")
    ws.append([f"账户：{acc.get('name', '')} ({acc.get('email', '')})"])
    ws.append(["时间", "类型", "金额(USDT)", "变动后余额", "关联单号/ID", "备注"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = hf, hfill, Alignment(horizontal="center")
    for r in rows:
        ws.append([(r.get("created_at") or "")[:19].replace("T", " "), LEDGER_CAT.get(r.get("category"), r.get("category")),
                   r.get("amount"), r.get("balance_after"), r.get("ref_id") or "", r.get("remark") or ""])
    for i, w in enumerate([20, 12, 14, 14, 24, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, len(rows), f"ledger_{acc.get('name', '')}_{month or start or 'all'}.xlsx"


# --- Background export jobs -------------------------------------------------
def _require_export(user: dict):
    if user["role"] == "operator" and "reports.export" not in (user.get("permissions") or []):
        raise HTTPException(status_code=403, detail="无报表导出权限")


def _export_notify_threshold(cfg: dict) -> int:
    try:
        return max(1, int(cfg.get("export_notify_threshold", 500) or 500))
    except (TypeError, ValueError):
        return 500


async def start_export_job(created_by: str, created_email: str, kind: str, allowed: dict, user_for_run: dict) -> str:
    doc = {"kind": kind, "status": "pending", "params": allowed,
           "created_by": created_by, "created_email": created_email, "created_at": iso(now_utc())}
    res = await db.export_jobs.insert_one(doc)
    jid = str(res.inserted_id)
    asyncio.create_task(_run_export(jid, user_for_run, kind, allowed))
    return jid


async def _run_export(job_id: str, user: dict, kind: str, params: dict):
    async def progress(done, total):
        await db.export_jobs.update_one({"_id": ObjectId(job_id)}, {"$set": {"processed": done, "expected": total}})
    try:
        if kind == "orders":
            buf, count, fname = await generate_orders_xlsx(user, progress=progress, **params)
        elif kind == "snapshots":
            buf, count, fname = await generate_snapshots_xlsx(**params)
        else:
            buf, count, fname = await generate_ledger_xlsx(user, **params)
        data = buf.getvalue()
        file_id = await fs.upload_from_stream(f"{job_id}.xlsx", data)
        cfg = await notify_settings()
        retain = max(1, int(cfg.get("export_retain_days", 7) or 7))
        await db.export_jobs.update_one({"_id": ObjectId(job_id)},
            {"$set": {"status": "done", "row_count": count, "filename": fname, "size_bytes": len(data),
                      "file_id": str(file_id), "processed": count, "expected": count,
                      "finished_at": iso(now_utc()), "expires_at": iso(now_utc() + timedelta(days=retain))}})
        if count >= _export_notify_threshold(cfg):
            kind_label = {"orders": "订单", "snapshots": "分润快照", "ledger": "资金流水"}.get(kind, kind)
            await send_telegram(f"📁 <b>导出完成</b>\n类型：{kind_label}\n行数：{count}\n大小：{round(len(data) / 1024)}KB\n操作人：{user.get('email', '')}", "report")
    except Exception as e:
        await db.export_jobs.update_one({"_id": ObjectId(job_id)},
            {"$set": {"status": "failed", "error": str(e)[:300], "finished_at": iso(now_utc())}})
        email = str(user.get("email", ""))
        if email.startswith("定时导出"):
            kind_label = {"orders": "订单", "snapshots": "分润快照", "ledger": "资金流水"}.get(kind, kind)
            # auto-retry once (retry job email carries ·重试 so it won't loop)
            if "重试" not in email:
                retry_user = {**user, "email": email + "·重试"}
                await start_export_job("system", retry_user["email"], kind, params, retry_user)
                await send_telegram(f"⚠️ <b>定时导出失败，正在自动重试</b>\n类型：{kind_label}\n原因：{str(e)[:180]}", "report")
            else:
                await send_telegram(f"❌ <b>定时导出重试仍失败</b>\n类型：{kind_label}\n原因：{str(e)[:180]}", "report")


@router.post("/export/{kind}")
async def create_export(kind: str, request: Request, user: dict = Depends(get_current_user)):
    _require_export(user)
    if kind not in EXPORT_KINDS:
        raise HTTPException(status_code=400, detail="不支持的导出类型")
    try:
        params = await request.json()
    except Exception:
        params = {}
    fields = {"orders": ("status", "network", "keyword", "merchant_id", "date", "start", "end", "ids"),
              "snapshots": ("start", "end"),
              "ledger": ("account_id", "month", "start", "end")}[kind]
    allowed = {k: params.get(k) for k in fields if params.get(k)}
    if kind == "snapshots" and user["role"] not in ("admin", "operator"):
        raise HTTPException(status_code=403, detail="无权限导出分润快照")
    if kind == "orders" and user["role"] in ("merchant", "agent"):
        scope = await merchant_ids_for_scope(user)
        if allowed.get("merchant_id") and (scope is None or ObjectId(allowed["merchant_id"]) not in scope):
            raise HTTPException(status_code=403, detail="无权限导出该商户订单")
    if kind == "ledger":
        if not allowed.get("account_id"):
            raise HTTPException(status_code=400, detail="导出资金流水需指定 account_id")
        allowed["account_id"] = str(await resolve_ledger_account(user, allowed["account_id"]))
    jid = await start_export_job(str(user["_id"]), user["email"], kind, allowed, user)
    await audit(user, "export.create", "export", jid, {"kind": kind, "params": allowed}, request)
    return {"id": jid, "status": "pending"}


@router.get("/export/jobs")
async def list_export_jobs(user: dict = Depends(get_current_user)):
    _require_export(user)
    q = {} if user["role"] == "admin" else {"created_by": str(user["_id"])}
    rows = await db.export_jobs.find(q).sort("created_at", -1).to_list(50)
    return [serialize(r) for r in rows]


@router.get("/export/jobs/{job_id}/download")
async def download_export_job(job_id: str, user: dict = Depends(get_current_user)):
    _require_export(user)
    job = await db.export_jobs.find_one({"_id": ObjectId(job_id)})
    if not job or job.get("status") != "done" or not job.get("file_id"):
        raise HTTPException(status_code=404, detail="文件未就绪")
    if user["role"] != "admin" and job.get("created_by") != str(user["_id"]):
        raise HTTPException(status_code=403, detail="无权下载")
    try:
        stream = await fs.open_download_stream(ObjectId(job["file_id"]))
        data = await stream.read()
    except Exception:
        raise HTTPException(status_code=410, detail="导出文件已过期，请重新导出")
    fname = job.get("filename", "orders.xlsx")
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# --- Scheduled automated exports -------------------------------------------
class ExportScheduleIn(BaseModel):
    kind: str
    freq: str  # daily | weekly
    range: str = "all"  # all | yesterday | week
    account_id: Optional[str] = None


@router.get("/export-schedules")
async def list_export_schedules(user: dict = Depends(require_roles("admin"))):
    rows = await db.export_schedules.find().sort("created_at", -1).to_list(100)
    out = []
    for r in rows:
        v = serialize(r)
        if r.get("account_id"):
            acc = await db.users.find_one({"_id": ObjectId(r["account_id"])})
            v["account_name"] = acc.get("name") if acc else r["account_id"]
        out.append(v)
    return out


@router.post("/export-schedules")
async def create_export_schedule(data: ExportScheduleIn, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.kind not in EXPORT_KINDS:
        raise HTTPException(status_code=400, detail="不支持的导出类型")
    if data.freq not in ("daily", "weekly"):
        raise HTTPException(status_code=400, detail="频率需为 daily 或 weekly")
    if data.range not in ("all", "yesterday", "week"):
        raise HTTPException(status_code=400, detail="区间需为 all / yesterday / week")
    if data.kind == "ledger" and not data.account_id:
        raise HTTPException(status_code=400, detail="资金流水定时导出需指定账户")
    doc = {"kind": data.kind, "freq": data.freq, "range": data.range, "account_id": data.account_id or None,
           "enabled": True, "created_by": str(user["_id"]), "created_email": user["email"],
           "created_at": iso(now_utc()), "last_run_at": None}
    res = await db.export_schedules.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit(user, "export.schedule.create", "export_schedule", res.inserted_id,
                {"kind": data.kind, "freq": data.freq}, request)
    return serialize(doc)


@router.patch("/export-schedules/{sid}/toggle")
async def toggle_export_schedule(sid: str, request: Request, user: dict = Depends(require_roles("admin"))):
    sch = await db.export_schedules.find_one({"_id": ObjectId(sid)})
    if not sch:
        raise HTTPException(status_code=404, detail="定时任务不存在")
    enabled = not sch.get("enabled", True)
    await db.export_schedules.update_one({"_id": sch["_id"]}, {"$set": {"enabled": enabled}})
    await audit(user, "export.schedule.toggle", "export_schedule", sid, {"enabled": enabled}, request)
    return {"ok": True, "enabled": enabled}


@router.delete("/export-schedules/{sid}")
async def delete_export_schedule(sid: str, request: Request, user: dict = Depends(require_roles("admin"))):
    res = await db.export_schedules.delete_one({"_id": ObjectId(sid)})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="定时任务不存在")
    await audit(user, "export.schedule.delete", "export_schedule", sid, {}, request)
    return {"ok": True}


async def run_scheduled_exports(run_id: str) -> dict:
    is_monday = now_utc().weekday() == 0
    sys_user = {"_id": "system", "email": "定时任务", "role": "admin", "permissions": ["reports.export"]}
    schedules = await db.export_schedules.find({"enabled": True}).to_list(200)
    started, skipped = 0, 0
    for sch in schedules:
        if sch.get("freq") == "weekly" and not is_monday:
            skipped += 1
            continue
        allowed = {}
        rng = sch.get("range", "all")
        if rng == "yesterday":
            y = (now_utc() - timedelta(days=1)).strftime("%Y-%m-%d")
            allowed["start"], allowed["end"] = y, y
        elif rng == "week":
            monday = now_utc() - timedelta(days=now_utc().weekday())
            allowed["start"], allowed["end"] = monday.strftime("%Y-%m-%d"), now_utc().strftime("%Y-%m-%d")
        if sch["kind"] == "ledger" and sch.get("account_id"):
            allowed["account_id"] = sch["account_id"]
        label = {"daily": "每日", "weekly": "每周"}.get(sch.get("freq"), sch.get("freq"))
        await start_export_job("system", f"定时导出·{label}", sch["kind"], allowed, sys_user)
        await db.export_schedules.update_one({"_id": sch["_id"]}, {"$set": {"last_run_at": iso(now_utc())}})
        started += 1
    logger.info(f"scheduled exports run {run_id}: started={started} skipped={skipped}")
    return {"started": started, "skipped": skipped}


@router.post("/cron/run-exports")
async def cron_run_exports(request: Request):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="unauthorized")
    try:
        body = await request.json()
    except Exception:
        body = {}
    run_id = request.headers.get("X-Webhook-Id") or (body or {}).get("run_id") or secrets.token_hex(8)
    asyncio.create_task(run_scheduled_exports(run_id))
    return {"ok": True, "accepted": True, "run_id": run_id}


# --- Export expiry reminders (24h ahead) -----------------------------------
async def run_export_expiry_reminders(run_id: str) -> dict:
    """Find completed export files expiring within the next 24h and push one
    Telegram reminder per file (once), so operators can re-download in time."""
    cfg = await notify_settings()
    if not cfg.get("export_expiry_notify", True):
        logger.info(f"export expiry reminders {run_id}: disabled")
        return {"notified": 0, "disabled": True}
    now = now_utc()
    soon = now + timedelta(hours=24)
    jobs = await db.export_jobs.find({
        "status": "done",
        "expiry_notified": {"$ne": True},
        "expires_at": {"$gte": iso(now), "$lt": iso(soon)},
    }).to_list(200)
    notified = 0
    kind_labels = {"orders": "订单", "snapshots": "分润快照", "ledger": "资金流水"}
    for j in jobs:
        try:
            exp = datetime.fromisoformat(j["expires_at"])
            hrs = max(0, int((exp - now).total_seconds() // 3600))
        except Exception:
            hrs = 0
        await send_telegram(
            f"⏳ <b>导出文件即将到期</b>\n类型：{kind_labels.get(j.get('kind'), j.get('kind'))}\n"
            f"文件：{j.get('filename', '')}\n行数：{j.get('row_count', '—')}\n"
            f"剩余约 {hrs} 小时后清理，请及时下载", "report")
        await db.export_jobs.update_one({"_id": j["_id"]}, {"$set": {"expiry_notified": True}})
        notified += 1
    logger.info(f"export expiry reminders {run_id}: notified={notified}")
    return {"notified": notified}


@router.post("/cron/export-expiry-reminders")
async def cron_export_expiry_reminders(request: Request):
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="unauthorized")
    try:
        body = await request.json()
    except Exception:
        body = {}
    run_id = request.headers.get("X-Webhook-Id") or (body or {}).get("run_id") or secrets.token_hex(8)
    asyncio.create_task(run_export_expiry_reminders(run_id))
    return {"ok": True, "accepted": True, "run_id": run_id}
