from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, Query
from fastapi.responses import JSONResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: F401  (client/db provided via core)
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Annotated, Any
from bson import ObjectId
from datetime import datetime, timezone, timedelta
import logging
import random
import string
import secrets
import bcrypt
import jwt
import asyncio
import hmac
import hashlib
import json
import re
import httpx
import pyotp
import qrcode
import io as _io
import base64
from email_utils import notify_security

# ---------------------------------------------------------------------------
# DB & app setup
# ---------------------------------------------------------------------------
from core import client, db, fs

app = FastAPI(title="四方收款平台 API")
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger("sifang")

JWT_ALGORITHM = "HS256"
NETWORKS = ["TRC20"]
TRONGRID_API = os.environ.get("TRONGRID_API", "https://api.trongrid.io")
TRONGRID_API_KEY = os.environ.get("TRONGRID_API_KEY", "")
USDT_TRC20_CONTRACT = "TR7NHqjeKQxGTCi8q8ZY4pL8otSzgjLj6t"
CHAIN_POLL_SECONDS = 20
CALLBACK_RETRY_DELAYS = [15, 60, 300, 1800, 7200]


# ---------------------------------------------------------------------------
# Helpers — pure/stateless helpers now live in utils.py
# ---------------------------------------------------------------------------
from utils import (
    now_utc, iso, hash_password, verify_password, get_jwt_secret, create_access_token,
    gen_order_no, gen_wallet_address, gen_txid, gen_api_key, gen_api_secret,
    clean_user, serialize, canon, sign_payload, client_ip,
)


async def audit(user: dict, action: str, target_type: str, target_id: str,
                detail: Any = None, request: Optional[Request] = None):
    await db.audit_logs.insert_one({
        "actor_id": str(user["_id"]), "actor_email": user["email"], "actor_role": user["role"],
        "action": action, "target_type": target_type, "target_id": str(target_id),
        "detail": detail or {}, "ip": client_ip(request) if request else None,
        "created_at": iso(now_utc()),
    })


async def record_ledger(account_id: str, category: str, amount: float, balance_after: float,
                        frozen_after: Optional[float] = None, ref_type: Optional[str] = None,
                        ref_id: Optional[Any] = None, remark: str = "", operator: Optional[str] = None):
    await db.balance_ledger.insert_one({
        "account_id": str(account_id), "category": category, "amount": round(amount, 4),
        "balance_after": round(balance_after, 4),
        "frozen_after": round(frozen_after, 4) if frozen_after is not None else None,
        "ref_type": ref_type, "ref_id": str(ref_id) if ref_id else None,
        "remark": remark, "operator": operator, "created_at": iso(now_utc()),
    })


async def enqueue_callback(order: dict):
    url = order.get("notify_url")
    if not url:
        return
    await db.callbacks.insert_one({
        "order_no": order["order_no"], "merchant_id": order["merchant_id"],
        "merchant_name": order["merchant_name"], "url": url, "status": "pending",
        "attempts": 0, "max_attempts": len(CALLBACK_RETRY_DELAYS) + 1,
        "next_at": iso(now_utc()), "last_code": None, "last_response": None,
        "created_at": iso(now_utc()), "updated_at": iso(now_utc()),
    })
    await db.orders.update_one({"_id": order["_id"]}, {"$set": {"callback_status": "pending"}})


async def build_callback_payload(order: dict, secret: str) -> dict:
    payload = {
        "order_no": order["order_no"], "amount_usdt": order["amount_usdt"],
        "pay_amount": order.get("pay_amount", order["amount_usdt"]), "network": order["network"],
        "status": order["status"], "tx_id": order.get("tx_id"), "paid_at": order.get("paid_at"),
        "remark": order.get("remark", ""), "out_trade_no": order.get("out_trade_no") or "",
        "timestamp": int(now_utc().timestamp()),
    }
    payload["sign"] = sign_payload(payload, secret)
    return payload


async def deliver_callback(cb: dict):
    order = await db.orders.find_one({"order_no": cb["order_no"]})
    merchant = await db.users.find_one({"_id": ObjectId(cb["merchant_id"])})
    if not order or not merchant:
        await db.callbacks.update_one({"_id": cb["_id"]}, {"$set": {"status": "failed", "last_response": "订单或商户不存在"}})
        return
    payload = await build_callback_payload(order, merchant.get("api_secret", ""))
    code, body, ok = None, "", False
    try:
        async with httpx.AsyncClient(timeout=10) as hc:
            resp = await hc.post(cb["url"], json=payload)
            code, body = resp.status_code, resp.text[:500]
            ok = 200 <= resp.status_code < 300 and ("success" in body.lower() or body.strip() in ("ok", "OK", "1", ""))
    except Exception as e:
        body = f"请求异常: {e}"[:500]
    attempts = cb["attempts"] + 1
    upd = {"attempts": attempts, "last_code": code, "last_response": body, "updated_at": iso(now_utc())}
    if ok:
        upd["status"] = "success"
        upd["delivered_at"] = iso(now_utc())
    elif attempts - 1 < len(CALLBACK_RETRY_DELAYS):
        upd["status"] = "retrying"
        upd["next_at"] = iso(now_utc() + timedelta(seconds=CALLBACK_RETRY_DELAYS[attempts - 1]))
    else:
        upd["status"] = "failed"
    await db.callbacks.update_one({"_id": cb["_id"]}, {"$set": upd})
    await db.orders.update_one({"_id": order["_id"]}, {"$set": {"callback_status": upd["status"]}})


async def callback_worker():
    while True:
        try:
            due = await db.callbacks.find({"status": {"$in": ["pending", "retrying"]},
                                           "next_at": {"$lte": iso(now_utc())}}).to_list(50)
            for cb in due:
                await deliver_callback(cb)
        except Exception as e:
            logger.warning(f"callback worker error: {e}")
        await asyncio.sleep(5)


async def fetch_trc20_transfers(address: str, since_ms: int) -> List[dict]:
    headers = {"TRON-PRO-API-KEY": TRONGRID_API_KEY} if TRONGRID_API_KEY else {}
    params = {"only_confirmed": "true", "only_to": "true", "limit": 100,
              "contract_address": USDT_TRC20_CONTRACT, "min_timestamp": since_ms}
    async with httpx.AsyncClient(timeout=15) as hc:
        r = await hc.get(f"{TRONGRID_API}/v1/accounts/{address}/transactions/trc20", params=params, headers=headers)
        r.raise_for_status()
        return [{"transaction_id": t["transaction_id"], "from": t.get("from"), "value": t["value"],
                 "decimals": t.get("token_info", {}).get("decimals", 6), "block_timestamp": t["block_timestamp"]}
                for t in r.json().get("data", [])]


async def fetch_transfers(network: str, address: str, since_ms: int) -> List[dict]:
    return await fetch_trc20_transfers(address, since_ms)


async def fetch_balance(network: str, address: str) -> float:
    headers = {"TRON-PRO-API-KEY": TRONGRID_API_KEY} if TRONGRID_API_KEY else {}
    async with httpx.AsyncClient(timeout=15) as hc:
        r = await hc.get(f"{TRONGRID_API}/v1/accounts/{address}", headers=headers)
        r.raise_for_status()
        data = r.json().get("data") or [{}]
    for item in data[0].get("trc20", []):
        if USDT_TRC20_CONTRACT in item:
            return round(int(item[USDT_TRC20_CONTRACT]) / 1e6, 2)
    return 0.0


DEFAULT_ALERT_SETTINGS = {"pct": 10.0, "min_abs": 100.0, "sync_minutes": 10, "sweep_threshold": 10000.0, "big_payment": 1000.0}
DEFAULT_NOTIFY_SETTINGS = {"telegram_enabled": False, "telegram_bot_token": "", "telegram_chat_id": "",
                           "notify_alerts": True, "notify_big_payment": True, "notify_sweep": True}


async def notify_settings() -> dict:
    doc = await db.settings.find_one({"_id": "notify"}) or {}
    doc.pop("_id", None)
    return {**DEFAULT_NOTIFY_SETTINGS, **doc}


async def send_telegram(text: str, kind: str = "alert") -> bool:
    cfg = await notify_settings()
    if not cfg["telegram_enabled"] or not cfg["telegram_bot_token"] or not cfg["telegram_chat_id"]:
        return False
    if kind != "report" and ((kind == "alert" and not cfg["notify_alerts"]) or (kind == "payment" and not cfg["notify_big_payment"]) or (kind == "sweep" and not cfg["notify_sweep"])):
        return False
    try:
        async with httpx.AsyncClient(timeout=10) as hc:
            r = await hc.post(f"https://api.telegram.org/bot{cfg['telegram_bot_token']}/sendMessage",
                              json={"chat_id": cfg["telegram_chat_id"], "text": text, "parse_mode": "HTML", "disable_web_page_preview": True})
            ok = r.status_code == 200 and r.json().get("ok")
        await db.notify_log.insert_one({"kind": kind, "text": text, "ok": bool(ok), "resp": r.text[:300], "created_at": iso(now_utc())})
        return bool(ok)
    except Exception as e:
        await db.notify_log.insert_one({"kind": kind, "text": text, "ok": False, "resp": str(e)[:300], "created_at": iso(now_utc())})
        return False


async def alert_settings() -> dict:
    doc = await db.settings.find_one({"_id": "alerts"}) or {}
    doc.pop("_id", None)
    return {**DEFAULT_ALERT_SETTINGS, **doc}


async def raise_alert(kind: str, network: str, address: str, channel: dict, message: str, detail: dict, cooldown_min: int = 60):
    recent = await db.alerts.find_one({"type": kind, "address": address, "network": network,
                                       "created_at": {"$gt": iso(now_utc() - timedelta(minutes=cooldown_min))}})
    if recent:
        return
    await db.alerts.insert_one({"type": kind, "network": network, "address": address, "channel_name": channel["name"],
                                "message": message, "detail": detail, "read": False, "created_at": iso(now_utc())})
    logger.warning(f"ALERT {kind} {network} {address}: {message}")
    icon = {"balance_change": "⚠️", "sync_failed": "🔴", "sweep_due": "🧊"}.get(kind, "⚠️")
    await send_telegram(f"{icon} <b>地址预警</b>\n{message}\n<code>{network} {address}</code>", "sweep" if kind == "sweep_due" else "alert")


async def check_sweep_alert(network: str, address: str, ch: dict, bal: float):
    cfg = await alert_settings()
    if cfg.get("sweep_threshold", 0) > 0 and bal >= cfg["sweep_threshold"]:
        await raise_alert("sweep_due", network, address, ch,
                          f"{ch['name']} 地址余额 {round(bal, 2)} USDT 已超过归集阈值 {cfg['sweep_threshold']}，建议归集至冷钱包",
                          {"balance": round(bal, 2), "threshold": cfg["sweep_threshold"]}, cooldown_min=360)


async def check_balance_alert(network: str, address: str, ch: dict, new_bal: float):
    prev = await db.address_health.find_one({"address": address, "network": network})
    old = prev.get("balance") if prev else None
    if old is None:
        return
    diff = new_bal - old
    cfg = await alert_settings()
    if abs(diff) >= cfg["min_abs"] and (old == 0 or abs(diff) / old * 100 >= cfg["pct"]):
        sign = "+" if diff > 0 else ""
        await raise_alert("balance_change", network, address, ch,
                          f"{ch['name']} 地址余额突变 {sign}{round(diff, 2)} USDT（{round(old, 2)} → {round(new_bal, 2)}）",
                          {"before": round(old, 2), "after": round(new_bal, 2), "diff": round(diff, 2)}, cooldown_min=5)


async def check_sync_alert(network: str, address: str, ch: dict):
    h = await db.address_health.find_one({"address": address, "network": network}) or {}
    cfg = await alert_settings()
    first_err = h.get("first_error_at")
    if not first_err:
        await db.address_health.update_one({"address": address, "network": network}, {"$set": {"first_error_at": iso(now_utc())}}, upsert=True)
        return
    if datetime.fromisoformat(first_err) <= now_utc() - timedelta(minutes=cfg["sync_minutes"]):
        await raise_alert("sync_failed", network, address, ch,
                          f"{ch['name']} 地址同步失败已超过 {cfg['sync_minutes']} 分钟：{h.get('sync_error', '')[:80]}",
                          {"since": first_err, "error": h.get("sync_error")})


async def set_address_health(network: str, address: str, channel: dict, **fields):
    fields["updated_at"] = iso(now_utc())
    await db.address_health.update_one({"address": address, "network": network},
                                       {"$set": {**fields, "channel_id": str(channel["_id"]), "channel_name": channel["name"]}},
                                       upsert=True)


async def match_chain_tx(tx: dict, address: str, network: str = "TRC20"):
    txid = tx["transaction_id"]
    if await db.chain_txs.find_one({"tx_id": txid}):
        return
    decimals = tx.get("decimals", 6)
    amount = round(int(tx["value"]) / (10 ** decimals), 6)
    ts = datetime.fromtimestamp(tx["block_timestamp"] / 1000, tz=timezone.utc)
    addr_q = address if network == "TRC20" else {"$regex": f"^{address}$", "$options": "i"}
    order = await db.orders.find_one({"status": "pending", "wallet_address": addr_q,
                                      "pay_amount": amount, "network": network})
    rec = {"tx_id": txid, "network": network, "to": address, "from": tx.get("from"),
           "amount": amount, "block_time": iso(ts), "order_no": order["order_no"] if order else None,
           "matched": bool(order), "created_at": iso(now_utc())}
    await db.chain_txs.insert_one(rec)
    if order:
        await mark_paid(order, tx_id=txid, paid_at=ts)
        logger.info(f"链上到账匹配成功 {order['order_no']} {amount} USDT tx={txid}")


def channel_addresses(ch: dict) -> List[str]:
    addrs = ch.get("wallet_addresses") or ([ch["wallet_address"]] if ch.get("wallet_address") else [])
    return [a for a in addrs if a]


async def pick_address(ch: dict) -> str:
    addrs = channel_addresses(ch)
    paused = {m["address"] async for m in db.address_meta.find({"address": {"$in": addrs}, "status": "paused"})}
    addrs = [a for a in addrs if a not in paused] or addrs
    if len(addrs) == 1:
        return addrs[0]
    pipeline = [{"$match": {"status": "pending", "wallet_address": {"$in": addrs}}},
                {"$group": {"_id": "$wallet_address", "n": {"$sum": 1}}}]
    counts = {r["_id"]: r["n"] async for r in db.orders.aggregate(pipeline)}
    return min(addrs, key=lambda a: (counts.get(a, 0), random.random()))


async def live_address_list() -> List[tuple]:
    chs = await db.channels.find({"status": "active"}).to_list(200)
    return [(ch, a) for ch in chs for a in channel_addresses(ch)]


async def sync_address(ch: dict, addr: str, since_ms: int) -> int:
    network = ch["network"]
    try:
        txs = await fetch_transfers(network, addr, since_ms)
        for tx in txs:
            await match_chain_tx(tx, addr, network)
        await db.channels.update_one({"_id": ch["_id"]}, {"$set": {"last_sync_at": iso(now_utc()), "sync_error": None}})
        await set_address_health(network, addr, ch, last_sync_at=iso(now_utc()), sync_error=None, last_tx_count=len(txs), first_error_at=None)
        return len(txs)
    except Exception as e:
        await db.channels.update_one({"_id": ch["_id"]}, {"$set": {"sync_error": str(e)[:200]}})
        await set_address_health(network, addr, ch, sync_error=str(e)[:200])
        await check_sync_alert(network, addr, ch)
        logger.warning(f"{network} 轮询失败 {addr}: {e}")
        raise


async def chain_listener():
    await asyncio.sleep(3)
    while True:
        try:
            for ch, addr in await live_address_list():
                has_pending = await db.orders.count_documents({"status": "pending", "wallet_address": addr}) > 0
                if not has_pending:
                    continue
                since = int((now_utc() - timedelta(hours=2)).timestamp() * 1000)
                try:
                    await sync_address(ch, addr, since)
                except Exception:
                    pass
                await asyncio.sleep(1)
            await db.orders.update_many({"status": "pending", "expires_at": {"$lt": iso(now_utc())}},
                                        {"$set": {"status": "expired"}})
        except Exception as e:
            logger.warning(f"chain listener error: {e}")
        await asyncio.sleep(CHAIN_POLL_SECONDS)


async def balance_refresher():
    await asyncio.sleep(10)
    while True:
        try:
            for ch, addr in await live_address_list():
                try:
                    bal = await fetch_balance(ch["network"], addr)
                    await check_balance_alert(ch["network"], addr, ch, bal)
                    await check_sweep_alert(ch["network"], addr, ch, bal)
                    await set_address_health(ch["network"], addr, ch, balance=bal, balance_at=iso(now_utc()), balance_error=None)
                except Exception as e:
                    await set_address_health(ch["network"], addr, ch, balance_error=str(e)[:200])
                await asyncio.sleep(2)
        except Exception as e:
            logger.warning(f"balance refresher error: {e}")
        await asyncio.sleep(300)


async def unique_pay_amount(amount: float, address: str) -> float:
    base = round(amount, 2)
    for _ in range(300):
        cents = random.randint(1, 9999) / 10000
        cand = round(base + cents, 4)
        exists = await db.orders.find_one({"status": "pending", "wallet_address": address, "pay_amount": cand})
        if not exists:
            return cand
    raise HTTPException(status_code=409, detail="该地址待支付订单过多，请稍后再试")


# ---------------------------------------------------------------------------
# Auth / permission dependencies — moved to deps.py
# ---------------------------------------------------------------------------
from deps import (
    get_current_user, require_roles, perm_for_path, merchant_ids_for_scope,
    OPERATOR_PERMS, PERM_RULES, OPERATOR_FORBIDDEN,
)


# ---------------------------------------------------------------------------
# Models — request models extracted to models.py
# ---------------------------------------------------------------------------
from models import (
    LoginInput, AccountCreate, PermissionsUpdate, BalanceAdjust, ChannelCreate, ChannelUpdate,
    OrderCreate, ApiOrderCreate, ApiOrderQuery, DevTestOrder, NotifySettings, IpWhitelist,
)


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
LOCK_MAX_FAILS = 5
LOCK_MINUTES = 15


async def check_lockout(identifier: str):
    rec = await db.login_attempts.find_one({"identifier": identifier})
    if rec and rec.get("locked_until") and datetime.fromisoformat(rec["locked_until"]) > now_utc():
        remain = int((datetime.fromisoformat(rec["locked_until"]) - now_utc()).total_seconds() // 60) + 1
        raise HTTPException(status_code=423, detail=f"失败次数过多，账号已临时锁定，请 {remain} 分钟后再试")


async def record_login_failure(identifier: str, email: str, ip: Optional[str]):
    rec = await db.login_attempts.find_one({"identifier": identifier})
    window_start = now_utc() - timedelta(minutes=LOCK_MINUTES)
    count = 1
    if rec and rec.get("last_fail_at") and datetime.fromisoformat(rec["last_fail_at"]) > window_start:
        count = rec.get("count", 0) + 1
    upd = {"identifier": identifier, "email": email, "ip": ip, "count": count, "last_fail_at": iso(now_utc())}
    locked = count >= LOCK_MAX_FAILS
    if locked:
        upd["locked_until"] = iso(now_utc() + timedelta(minutes=LOCK_MINUTES))
    await db.login_attempts.update_one({"identifier": identifier}, {"$set": upd}, upsert=True)
    await db.audit_logs.insert_one({
        "actor_id": None, "actor_email": email, "actor_role": "anonymous",
        "action": "auth.locked" if locked else "auth.login_failed", "target_type": "login", "target_id": email,
        "detail": {"fail_count": count, **({"locked_minutes": LOCK_MINUTES} if locked else {})},
        "ip": ip, "created_at": iso(now_utc()),
    })
    return locked


# ---------------------------------------------------------------------------
# Auth routes (login / logout / me / 2FA) moved to routes/auth.py
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Security policy
# ---------------------------------------------------------------------------
class SecurityPolicy(BaseModel):
    enforce_admin_2fa: bool = False
    enforce_operator_2fa: bool = False
    login_alert: bool = True


async def security_policy() -> dict:
    doc = await db.settings.find_one({"_id": "security"}) or {}
    doc.pop("_id", None)
    return {"enforce_admin_2fa": False, "enforce_operator_2fa": False, "login_alert": True, **doc}


@api_router.get("/security/policy")
async def get_security_policy(user: dict = Depends(get_current_user)):
    pol = await security_policy()
    enforced = (user["role"] == "admin" and pol["enforce_admin_2fa"]) or (user["role"] == "operator" and pol["enforce_operator_2fa"])
    return {**pol, "enforced_for_me": bool(enforced)}


@api_router.put("/security/policy")
async def put_security_policy(data: SecurityPolicy, request: Request, user: dict = Depends(require_roles("admin"))):
    await db.settings.update_one({"_id": "security"}, {"$set": data.model_dump()}, upsert=True)
    await audit(user, "settings.security", "settings", "security", data.model_dump(), request)
    return {"ok": True, **data.model_dump()}


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@api_router.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user), date: Optional[str] = None):
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}

    today_start = now_utc().replace(hour=0, minute=0, second=0, microsecond=0)
    if date:
        try:
            today_start = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="date 格式应为 YYYY-MM-DD")
    day_end = today_start + timedelta(days=1)
    all_orders = await db.orders.find(q).to_list(5000)

    total_volume = sum(o["amount_usdt"] for o in all_orders if o["status"] == "paid")
    today_volume = sum(o["amount_usdt"] for o in all_orders
                       if o["status"] == "paid" and today_start <= datetime.fromisoformat(o["created_at"]) < day_end)
    total_orders = len(all_orders)
    paid_orders = len([o for o in all_orders if o["status"] == "paid"])
    pending_orders = len([o for o in all_orders if o["status"] == "pending"])
    success_rate = round(paid_orders / total_orders * 100, 1) if total_orders else 0.0

    today_all = [o for o in all_orders if today_start <= datetime.fromisoformat(o["created_at"]) < day_end]
    today_orders = len(today_all)
    today_paid = len([o for o in today_all if o["status"] == "paid"])
    today_success_rate = round(today_paid / today_orders * 100, 1) if today_orders else 0.0

    if user["role"] == "admin":
        profit = sum(o.get("platform_profit", 0) for o in all_orders if o["status"] == "paid")
        profit_label = "平台利润"
    elif user["role"] == "agent":
        profit = sum(o.get("agent_commission", 0) for o in all_orders if o["status"] == "paid")
        profit_label = "分润收益"
    else:
        profit = sum(o.get("merchant_receive", 0) for o in all_orders if o["status"] == "paid")
        profit_label = "实收金额"

    # 7-day trend
    trend = []
    for i in range(6, -1, -1):
        day = (today_start - timedelta(days=i))
        day_end = day + timedelta(days=1)
        vol = sum(o["amount_usdt"] for o in all_orders
                  if o["status"] == "paid" and day <= datetime.fromisoformat(o["created_at"]) < day_end)
        cnt = len([o for o in all_orders if day <= datetime.fromisoformat(o["created_at"]) < day_end])
        trend.append({"date": day.strftime("%m-%d"), "volume": round(vol, 2), "orders": cnt})

    active_channels = await db.channels.count_documents({"status": "active"})
    balance = round(user.get("balance", 0), 2)

    return {
        "total_volume": round(total_volume, 2),
        "today_volume": round(today_volume, 2),
        "total_orders": total_orders,
        "pending_orders": pending_orders,
        "success_rate": success_rate,
        "profit": round(profit, 2),
        "profit_label": profit_label,
        "active_channels": active_channels,
        "balance": balance,
        "trend": trend,
        "today_orders": today_orders,
        "today_paid": today_paid,
        "today_success_rate": today_success_rate,
        "stat_date": today_start.strftime("%Y-%m-%d"),
    }


def _date_window(date: Optional[str], start: Optional[str], end: Optional[str]):
    if start and end:
        try:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
        if de <= ds:
            raise HTTPException(status_code=400, detail="结束日期需不早于开始日期")
        return ds, de, f"{start} ~ {end}"
    ds = now_utc().replace(hour=0, minute=0, second=0, microsecond=0)
    if date:
        try:
            ds = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="date 格式应为 YYYY-MM-DD")
    return ds, ds + timedelta(days=1), ds.strftime("%Y-%m-%d")


@api_router.get("/stats/merchants")
async def merchant_stats(date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                         user: dict = Depends(require_roles("admin"))):
    day_start, day_end, label = _date_window(date, start, end)
    merchants = await db.users.find({"role": "merchant"}).to_list(1000)
    orders = await db.orders.find({"created_at": {"$gte": iso(day_start), "$lt": iso(day_end)}}).to_list(100000)
    by_m: dict = {}
    for o in orders:
        s = by_m.setdefault(o["merchant_id"], {"orders": 0, "paid": 0, "volume": 0.0, "fee": 0.0, "receive": 0.0})
        s["orders"] += 1
        if o["status"] == "paid":
            s["paid"] += 1
            s["volume"] += o.get("amount_usdt", 0)
            s["fee"] += o.get("fee", 0)
            s["receive"] += o.get("merchant_receive", 0)
    result = []
    for mer in merchants:
        mid = str(mer["_id"])
        s = by_m.get(mid, {"orders": 0, "paid": 0, "volume": 0.0, "fee": 0.0, "receive": 0.0})
        result.append({
            "id": mid, "name": mer.get("name"), "email": mer.get("email"),
            "status": mer.get("status", "active"),
            "balance": round(mer.get("balance", 0), 2), "frozen_balance": round(mer.get("frozen_balance", 0), 2),
            "orders": s["orders"], "paid": s["paid"],
            "success_rate": round(s["paid"] / s["orders"] * 100, 1) if s["orders"] else 0.0,
            "volume": round(s["volume"], 2), "fee": round(s["fee"], 2), "receive": round(s["receive"], 2),
        })
    result.sort(key=lambda x: (x["volume"], x["paid"]), reverse=True)
    return {"date": label, "start": day_start.strftime("%Y-%m-%d"),
            "end": (day_end - timedelta(days=1)).strftime("%Y-%m-%d"),
            "count": len(result), "merchants": result}


@api_router.get("/reports/merchant-stats")
async def merchant_stats_report(date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                                user: dict = Depends(require_roles("admin"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    import io
    data = await merchant_stats(date=date, start=start, end=end, user=user)
    wb = Workbook(); ws = wb.active; ws.title = "商户统计"
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws.append([f"统计区间：{data['date']}"])
    ws.append(["商户", "邮箱", "订单数", "成功", "成功率%", "收款金额", "手续费", "实收", "可用余额"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for m in data["merchants"]:
        ws.append([m["name"], m["email"], m["orders"], m["paid"], m["success_rate"],
                   m["volume"], m["fee"], m["receive"], m["balance"]])
    for i, w in enumerate([18, 26, 10, 8, 10, 16, 14, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"merchant_stats_{data['date'].replace(' ', '').replace('~', '-')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------------------------------------------------------------------------
# Orders
# ---------------------------------------------------------------------------
async def pick_channel(network: Optional[str], amount: float) -> Optional[dict]:
    q = {"status": "active"}
    if network:
        q["network"] = network
    channels = await db.channels.find(q).to_list(100)
    channels = [c for c in channels if c["min_amount"] <= amount <= c["max_amount"]]
    if not channels:
        return None
    # weighted random
    total = sum(c["weight"] for c in channels)
    r = random.uniform(0, total)
    upto = 0
    for c in channels:
        upto += c["weight"]
        if r <= upto:
            return c
    return channels[0]


async def build_order(merchant: dict, amount_usdt: float, network: Optional[str], remark: str,
                      notify_url: Optional[str], out_trade_no: Optional[str] = None) -> dict:
    network = "TRC20"
    if amount_usdt <= 0:
        raise HTTPException(status_code=400, detail="金额必须大于 0")
    channel = await pick_channel(network, amount_usdt)
    if not channel:
        raise HTTPException(status_code=400, detail="没有可用的收款通道")
    expires = now_utc() + timedelta(minutes=15)
    live = bool(channel_addresses(channel))
    address = await pick_address(channel) if live else gen_wallet_address(network)
    amount = round(amount_usdt, 2)
    pay_amount = await unique_pay_amount(amount, address) if live else amount
    order = {
        "order_no": gen_order_no(),
        "out_trade_no": out_trade_no or None,
        "merchant_id": str(merchant["_id"]),
        "merchant_name": merchant["name"],
        "amount_usdt": amount,
        "pay_amount": pay_amount,
        "live": live,
        "network": network,
        "wallet_address": address,
        "tx_id": None,
        "status": "pending",
        "channel_id": str(channel["_id"]),
        "channel_name": channel["name"],
        "remark": remark or "",
        "notify_url": notify_url or merchant.get("notify_url") or None,
        "callback_status": None,
        "created_at": iso(now_utc()),
        "paid_at": None,
        "expires_at": iso(expires),
    }
    res = await db.orders.insert_one(order)
    order["_id"] = res.inserted_id
    return order


@api_router.post("/orders")
async def create_order(data: OrderCreate, user: dict = Depends(get_current_user)):
    # determine merchant
    if user["role"] == "merchant":
        merchant = user
    else:
        if not data.merchant_id:
            raise HTTPException(status_code=400, detail="请选择商户")
        merchant = await db.users.find_one({"_id": ObjectId(data.merchant_id), "role": "merchant"})
        if not merchant:
            raise HTTPException(status_code=404, detail="商户不存在")
    order = await build_order(merchant, data.amount_usdt, data.network, data.remark, data.notify_url, data.out_trade_no)
    return serialize(order)


# ---------------------------------------------------------------------------
# Merchant Open API (api_key + HMAC signature)
# ---------------------------------------------------------------------------
def api_pay_url(order_no: str) -> str:
    return f"{os.environ.get('FRONTEND_URL', '')}/pay/{order_no}"


async def verify_api_request(body: dict, request: Request) -> dict:
    merchant = await db.users.find_one({"api_key": body.get("api_key"), "role": "merchant"})
    if not merchant:
        raise HTTPException(status_code=401, detail="api_key 无效")
    if merchant.get("status") == "frozen":
        raise HTTPException(status_code=403, detail="商户已被冻结")
    wl = merchant.get("ip_whitelist") or []
    ip = client_ip(request)
    if wl and ip not in wl:
        await db.audit_logs.insert_one({
            "actor_id": str(merchant["_id"]), "actor_email": merchant["email"], "actor_role": "merchant",
            "action": "api.ip_rejected", "target_type": "api", "target_id": body.get("api_key", ""),
            "detail": {"ip": ip}, "ip": ip, "created_at": iso(now_utc())})
        raise HTTPException(status_code=403, detail=f"来源 IP {ip} 不在白名单内")
    if abs(int(now_utc().timestamp()) - int(body.get("timestamp", 0))) > 300:
        raise HTTPException(status_code=401, detail="timestamp 已过期（允许 ±300 秒）")
    nonce = body.get("nonce", "")
    if not nonce or await db.api_nonces.find_one({"api_key": body["api_key"], "nonce": nonce}):
        raise HTTPException(status_code=401, detail="nonce 重复或缺失")
    expected = sign_payload(body, merchant.get("api_secret", ""))
    if not hmac.compare_digest(expected, str(body.get("sign", ""))):
        raise HTTPException(status_code=401, detail="签名错误")
    await db.api_nonces.insert_one({"api_key": body["api_key"], "nonce": nonce, "created_at": now_utc()})
    return merchant


def api_order_view(o: dict) -> dict:
    return {"order_no": o["order_no"], "out_trade_no": o.get("out_trade_no"), "amount_usdt": o["amount_usdt"],
            "pay_amount": o.get("pay_amount", o["amount_usdt"]), "network": o["network"],
            "wallet_address": o["wallet_address"], "status": o["status"], "tx_id": o.get("tx_id"),
            "paid_at": o.get("paid_at"), "expires_at": o["expires_at"], "pay_url": api_pay_url(o["order_no"])}


@api_router.post("/v1/orders")
async def api_create_order(data: ApiOrderCreate, request: Request):
    body = data.model_dump()
    merchant = await verify_api_request(body, request)
    dup = await db.orders.find_one({"merchant_id": str(merchant["_id"]), "out_trade_no": data.out_trade_no})
    if dup:
        return {"code": 0, "msg": "duplicate", "data": api_order_view(dup)}
    order = await build_order(merchant, data.amount_usdt, data.network, data.remark, data.notify_url, data.out_trade_no)
    return {"code": 0, "msg": "success", "data": api_order_view(order)}


@api_router.post("/v1/orders/query")
async def api_query_order(data: ApiOrderQuery, request: Request):
    body = data.model_dump()
    merchant = await verify_api_request(body, request)
    q: dict = {"merchant_id": str(merchant["_id"])}
    if data.order_no:
        q["order_no"] = data.order_no
    elif data.out_trade_no:
        q["out_trade_no"] = data.out_trade_no
    else:
        raise HTTPException(status_code=400, detail="需提供 order_no 或 out_trade_no")
    o = await db.orders.find_one(q)
    if not o:
        raise HTTPException(status_code=404, detail="订单不存在")
    return {"code": 0, "msg": "success", "data": api_order_view(o)}


# ---------------------------------------------------------------------------
# Developer API console routes extracted to routes/dev.py (included at bottom)
# ---------------------------------------------------------------------------


@api_router.get("/notifications/paid")
async def paid_notifications(since: str, user: dict = Depends(get_current_user)):
    scope = await merchant_ids_for_scope(user)
    q: dict = {"status": "paid", "paid_at": {"$gt": since}}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    orders = await db.orders.find(q).sort("paid_at", -1).to_list(20)
    return {"now": iso(now_utc()), "orders": [
        {"order_no": o["order_no"], "amount_usdt": o["amount_usdt"], "merchant_name": o["merchant_name"],
         "network": o["network"], "live": o.get("live", False), "paid_at": o["paid_at"]} for o in orders]}


@api_router.get("/orders")
async def list_orders(user: dict = Depends(get_current_user),
                      status: Optional[str] = None, network: Optional[str] = None,
                      keyword: Optional[str] = None, merchant_id: Optional[str] = None,
                      date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                      limit: int = Query(200, le=1000)):
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    if merchant_id and (scope is None or ObjectId(merchant_id) in scope):
        q["merchant_id"] = str(merchant_id)
    if start and end:
        try:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    elif date:
        try:
            ds = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="date 格式应为 YYYY-MM-DD")
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(ds + timedelta(days=1))}
    if status:
        q["status"] = status
    if network:
        q["network"] = network
    if keyword:
        kw = keyword.strip()
        q["$or"] = [{"order_no": {"$regex": kw, "$options": "i"}},
                    {"tx_id": {"$regex": kw, "$options": "i"}},
                    {"out_trade_no": {"$regex": kw, "$options": "i"}},
                    {"wallet_address": {"$regex": kw, "$options": "i"}},
                    {"merchant_name": {"$regex": kw, "$options": "i"}}]
        if re.fullmatch(r"\d+(\.\d+)?", kw):
            q["$or"] += [{"amount_usdt": float(kw)}, {"pay_amount": float(kw)}]
    orders = await db.orders.find(q).sort("created_at", -1).to_list(limit)
    return [serialize(o) for o in orders]


@api_router.get("/orders/{order_no}/public")
async def get_order_public(order_no: str):
    order = await db.orders.find_one({"order_no": order_no})
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    # auto-expire
    if order["status"] == "pending" and datetime.fromisoformat(order["expires_at"]) < now_utc():
        await db.orders.update_one({"_id": order["_id"]}, {"$set": {"status": "expired"}})
        order["status"] = "expired"
    net = order["network"]
    explorer = f"https://tronscan.org/#/address/{order['wallet_address']}"
    tx_url = None
    if order.get("tx_id") and order.get("live"):
        tx_url = f"https://tronscan.org/#/transaction/{order['tx_id']}"
    return {
        "order_no": order["order_no"],
        "amount_usdt": order["amount_usdt"],
        "pay_amount": order.get("pay_amount", order["amount_usdt"]),
        "live": order.get("live", False),
        "network": net,
        "wallet_address": order["wallet_address"],
        "status": order["status"],
        "merchant_name": order["merchant_name"],
        "created_at": order["created_at"],
        "expires_at": order["expires_at"],
        "tx_id": order.get("tx_id"),
        "explorer_url": explorer,
        "tx_url": tx_url,
    }


DEFAULT_CHECKOUT_SETTINGS = {
    "energy_enabled": False, "energy_trx_address": "", "energy_price_text": "2TRX/笔（即每65500能量只需2TRX）",
    "support_url": "", "tutorial_links": [], "energy_rules": [],
    "logo_url": "", "brand_subtitle": "USDT 安全收款", "brand_subtitle_tw": "USDT 安全收款", "brand_subtitle_en": "Secure USDT Payment",
    "theme_color": "#26A17B",
}


@api_router.get("/checkout/settings")
async def get_checkout_settings():
    doc = await db.settings.find_one({"_id": "checkout"}) or {}
    doc.pop("_id", None)
    return {**DEFAULT_CHECKOUT_SETTINGS, **doc}


class CheckoutSettings(BaseModel):
    energy_enabled: bool = False
    energy_trx_address: str = ""
    energy_price_text: str = ""
    support_url: str = ""
    tutorial_links: List[dict] = []
    energy_rules: List[dict] = []
    logo_url: str = ""
    brand_subtitle: str = ""
    brand_subtitle_tw: str = ""
    brand_subtitle_en: str = ""
    theme_color: str = "#26A17B"


@api_router.put("/checkout/settings")
async def put_checkout_settings(data: CheckoutSettings, request: Request, user: dict = Depends(require_roles("admin"))):
    addr = data.energy_trx_address.strip()
    if addr:
        validate_wallet("TRC20", addr)
    import re
    if not re.fullmatch(r"#[0-9a-fA-F]{6}", data.theme_color or ""):
        raise HTTPException(status_code=400, detail="主题色需为 #RRGGBB 格式")
    doc = data.model_dump()
    doc["energy_trx_address"] = addr
    await db.settings.update_one({"_id": "checkout"}, {"$set": doc}, upsert=True)
    await audit(user, "settings.checkout", "settings", "checkout", doc, request)
    return doc


async def settle_order(order: dict):
    """Compute fee/commission split and credit balances."""
    merchant = await db.users.find_one({"_id": ObjectId(order["merchant_id"])})
    channel = await db.channels.find_one({"_id": ObjectId(order["channel_id"])})
    amount = order["amount_usdt"]
    merchant_rate = merchant.get("rate", 0.03)
    channel_rate = channel.get("fee_rate", 0.02) if channel else 0.02

    fee = round(amount * merchant_rate, 4)
    merchant_receive = round(amount - fee, 4)

    agent_commission = 0.0
    agent_id = merchant.get("parent_agent_id")
    if agent_id:
        agent = await db.users.find_one({"_id": ObjectId(agent_id)})
        if agent:
            agent_rate = agent.get("rate", channel_rate)
            agent_commission = round(amount * max(merchant_rate - agent_rate, 0), 4)
            if agent_commission > 0:
                await db.users.update_one({"_id": agent["_id"]},
                                          {"$inc": {"balance": agent_commission}})
                agent_upd = await db.users.find_one({"_id": agent["_id"]})
                await record_ledger(str(agent["_id"]), "commission", agent_commission,
                                    agent_upd.get("balance", 0), agent_upd.get("frozen_balance", 0),
                                    "order", order.get("order_no"),
                                    remark=f"订单 {order.get('order_no')} 代理分润")
    platform_profit = round(fee - agent_commission, 4)

    await db.users.update_one({"_id": merchant["_id"]}, {"$inc": {"balance": merchant_receive}})
    m_upd = await db.users.find_one({"_id": merchant["_id"]})
    await record_ledger(str(merchant["_id"]), "order_credit", merchant_receive,
                        m_upd.get("balance", 0), m_upd.get("frozen_balance", 0),
                        "order", order.get("order_no"),
                        remark=f"{order.get('network', '')} 收款 {amount} USDT，手续费 {fee}")
    return fee, merchant_receive, agent_commission, platform_profit


async def mark_paid(order: dict, tx_id: Optional[str] = None, paid_at: Optional[datetime] = None,
                    notify: bool = True) -> dict:
    claimed = await db.orders.find_one_and_update(
        {"_id": order["_id"], "status": "pending"}, {"$set": {"status": "confirming"}})
    if not claimed:
        return order
    fee, merchant_receive, agent_commission, platform_profit = await settle_order(order)
    update = {
        "status": "paid",
        "tx_id": tx_id or gen_txid(order["network"]),
        "paid_at": iso(paid_at or now_utc()),
        "fee": fee,
        "merchant_receive": merchant_receive,
        "agent_commission": agent_commission,
        "platform_profit": platform_profit,
        "callback_status": "none" if not order.get("notify_url") else "pending",
    }
    await db.orders.update_one({"_id": order["_id"]}, {"$set": update})
    order.update(update)
    if notify:
        await enqueue_callback(order)
    cfg = await alert_settings()
    if order["amount_usdt"] >= cfg.get("big_payment", 1000):
        await send_telegram(f"💰 <b>大额到账</b> {order['amount_usdt']} USDT\n商户：{order['merchant_name']}\n订单：<code>{order['order_no']}</code>\n网络：{order['network']}" + (f"\nTx：<code>{update['tx_id']}</code>" if order.get("live") else ""), "payment")
    return order


@api_router.post("/orders/{order_no}/simulate-pay")
async def simulate_pay(order_no: str):
    order = await db.orders.find_one({"order_no": order_no})
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    if order["status"] == "paid":
        return {"ok": True, "status": "paid"}
    if order["status"] != "pending":
        raise HTTPException(status_code=400, detail="订单状态不可支付")
    if order.get("live"):
        raise HTTPException(status_code=400, detail="真实链上订单不支持模拟支付，请向收款地址转账")
    await mark_paid(order)
    return {"ok": True, "status": "paid"}


@api_router.post("/orders/{order_id}/verify")
async def manual_verify(order_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    order = await db.orders.find_one({"_id": ObjectId(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    if order["status"] == "paid":
        raise HTTPException(status_code=400, detail="订单已支付")
    await db.orders.update_one({"_id": order["_id"]}, {"$set": {"status": "pending"}})
    order["status"] = "pending"
    await mark_paid(order)
    await audit(user, "order.manual_verify", "order", order["order_no"],
                {"amount_usdt": order["amount_usdt"]}, request)
    return {"ok": True}


@api_router.post("/orders/{order_id}/reject")
async def reject_order(order_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    order = await db.orders.find_one({"_id": ObjectId(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    if order["status"] == "paid":
        raise HTTPException(status_code=400, detail="订单已支付，无法驳回")
    if order["status"] == "rejected":
        return {"ok": True, "status": "rejected"}
    await db.orders.update_one({"_id": order["_id"]}, {"$set": {"status": "rejected", "rejected_at": iso(now_utc())}})
    await audit(user, "order.reject", "order", order["order_no"], {"amount_usdt": order["amount_usdt"]}, request)
    return {"ok": True, "status": "rejected"}


@api_router.post("/orders/{order_id}/resend-callback")
async def resend_callback(order_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    order = await db.orders.find_one({"_id": ObjectId(order_id)})
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    if order["status"] != "paid":
        raise HTTPException(status_code=400, detail="仅已支付订单可手动回调")
    if not order.get("notify_url"):
        raise HTTPException(status_code=400, detail="该订单未配置回调地址")
    await enqueue_callback(order)
    cb = await db.callbacks.find_one({"order_no": order["order_no"]}, sort=[("created_at", -1)])
    if cb:
        await deliver_callback(cb)
        cb = await db.callbacks.find_one({"_id": cb["_id"]})
    await audit(user, "order.resend_callback", "order", order["order_no"], {}, request)
    return {"ok": True, "callback": serialize(cb) if cb else None}


# ---------------------------------------------------------------------------
# Batch order operations (admin / operator)
# ---------------------------------------------------------------------------
class BatchOrderIds(BaseModel):
    ids: List[str]


@api_router.post("/orders/batch-reject")
async def batch_reject_orders(data: BatchOrderIds, request: Request, user: dict = Depends(require_roles("admin"))):
    success, failed = [], []
    for oid in data.ids:
        try:
            order = await db.orders.find_one({"_id": ObjectId(oid)})
            if not order:
                failed.append({"id": oid, "reason": "订单不存在"}); continue
            if order["status"] == "paid":
                failed.append({"id": oid, "order_no": order["order_no"], "reason": "已支付，无法驳回"}); continue
            if order["status"] == "rejected":
                success.append({"id": oid, "order_no": order["order_no"]}); continue
            await db.orders.update_one({"_id": order["_id"]}, {"$set": {"status": "rejected", "rejected_at": iso(now_utc())}})
            await audit(user, "order.reject", "order", order["order_no"], {"amount_usdt": order["amount_usdt"], "batch": True}, request)
            success.append({"id": oid, "order_no": order["order_no"]})
        except Exception as e:
            failed.append({"id": oid, "reason": str(e)[:100]})
    return {"ok": True, "total": len(data.ids), "success_count": len(success), "failed_count": len(failed),
            "success": success, "failed": failed}


@api_router.post("/orders/batch-resend-callback")
async def batch_resend_callback(data: BatchOrderIds, request: Request, user: dict = Depends(require_roles("admin"))):
    success, failed = [], []
    for oid in data.ids:
        try:
            order = await db.orders.find_one({"_id": ObjectId(oid)})
            if not order:
                failed.append({"id": oid, "reason": "订单不存在"}); continue
            if order["status"] != "paid":
                failed.append({"id": oid, "order_no": order["order_no"], "reason": "仅已支付订单可回调"}); continue
            if not order.get("notify_url"):
                failed.append({"id": oid, "order_no": order["order_no"], "reason": "未配置回调地址"}); continue
            await enqueue_callback(order)
            cb = await db.callbacks.find_one({"order_no": order["order_no"]}, sort=[("created_at", -1)])
            if cb:
                await deliver_callback(cb)
            await audit(user, "order.resend_callback", "order", order["order_no"], {"batch": True}, request)
            success.append({"id": oid, "order_no": order["order_no"]})
        except Exception as e:
            failed.append({"id": oid, "reason": str(e)[:100]})
    return {"ok": True, "total": len(data.ids), "success_count": len(success), "failed_count": len(failed),
            "success": success, "failed": failed}


# ---------------------------------------------------------------------------
# Daily profit-sharing reconciliation (cron): snapshot + backfill + Telegram
# Profit-sharing stays real-time (settle_order). This only snapshots & safety-nets.
# ---------------------------------------------------------------------------
async def build_profit_snapshot(day: str) -> dict:
    start, end = f"{day}T00:00:00", f"{day}T23:59:59.999999"
    orders = await db.orders.find({"status": "paid", "paid_at": {"$gte": start, "$lte": end}}).to_list(100000)
    merchants: dict = {}
    tot = {"amount": 0.0, "fee": 0.0, "receive": 0.0, "commission": 0.0, "platform": 0.0}
    for o in orders:
        tot["amount"] += o.get("amount_usdt", 0); tot["fee"] += o.get("fee", 0)
        tot["receive"] += o.get("merchant_receive", 0); tot["commission"] += o.get("agent_commission", 0)
        tot["platform"] += o.get("platform_profit", 0)
        m = merchants.setdefault(o["merchant_id"], {"merchant_id": o["merchant_id"], "merchant_name": o.get("merchant_name", ""),
                                                     "orders": 0, "amount": 0.0, "fee": 0.0, "receive": 0.0, "commission": 0.0, "platform": 0.0})
        m["orders"] += 1; m["amount"] += o.get("amount_usdt", 0); m["fee"] += o.get("fee", 0)
        m["receive"] += o.get("merchant_receive", 0); m["commission"] += o.get("agent_commission", 0); m["platform"] += o.get("platform_profit", 0)
    return {
        "date": day, "order_count": len(orders),
        "total_amount": round(tot["amount"], 4), "total_fee": round(tot["fee"], 4),
        "total_merchant_receive": round(tot["receive"], 4), "total_agent_commission": round(tot["commission"], 4),
        "total_platform_profit": round(tot["platform"], 4),
        "merchants": [{k: (round(v, 4) if isinstance(v, float) else v) for k, v in m.items()} for m in merchants.values()],
        "updated_at": iso(now_utc()),
    }


async def backfill_unsettled() -> int:
    count = 0
    async for o in db.orders.find({"status": "paid", "merchant_receive": None}):
        fee, mr, ac, pp = await settle_order(o)
        await db.orders.update_one({"_id": o["_id"]},
                                   {"$set": {"fee": fee, "merchant_receive": mr, "agent_commission": ac, "platform_profit": pp}})
        count += 1
    return count


def snapshot_report_text(snap: dict) -> str:
    bf = snap.get("backfilled") or 0
    return (f"📊 <b>日终分润对账</b> {snap.get('date')}\n"
            f"订单数：{snap['order_count']}\n收款总额：{snap['total_amount']} USDT\n"
            f"手续费：{snap['total_fee']} USDT\n商户实收：{snap['total_merchant_receive']} USDT\n"
            f"代理分润：{snap['total_agent_commission']} USDT\n平台利润：{snap['total_platform_profit']} USDT"
            + (f"\n⚠️ 补偿结算：{bf} 单" if bf else ""))


async def run_daily_settlement(run_id: str, day: str) -> dict:
    if await db.cron_runs.find_one({"run_id": run_id}):
        return {"skipped": True}
    await db.cron_runs.insert_one({"run_id": run_id, "job": "daily-settle", "day": day, "started_at": iso(now_utc())})
    backfilled = await backfill_unsettled()
    snap = await build_profit_snapshot(day)
    snap["backfilled"] = backfilled
    prev_day = (datetime.strptime(day, "%Y-%m-%d") - timedelta(days=1)).strftime("%Y-%m-%d")
    prev = await db.profit_snapshots.find_one({"date": prev_day})
    pct = None
    if prev and prev.get("total_platform_profit"):
        pct = round((snap["total_platform_profit"] - prev["total_platform_profit"]) / prev["total_platform_profit"] * 100, 1)
    snap["platform_profit_change_pct"] = pct
    await db.profit_snapshots.update_one({"date": day}, {"$set": snap}, upsert=True)
    text = snapshot_report_text(snap)
    if pct is not None and abs(pct) >= 50:
        text += f"\n🚨 平台利润环比{'骤增' if pct > 0 else '骤降'} {abs(pct)}%（异常波动预警）"
    await send_telegram(text, "report")
    await db.cron_runs.update_one({"run_id": run_id}, {"$set": {"finished_at": iso(now_utc())}})
    # cleanup expired export files (TTL)
    now_iso = iso(now_utc())
    async for ej in db.export_jobs.find({"expires_at": {"$lt": now_iso}, "file_id": {"$ne": None}}):
        try:
            await fs.delete(ObjectId(ej["file_id"]))
        except Exception:
            pass
        await db.export_jobs.delete_one({"_id": ej["_id"]})
    return snap


@api_router.post("/cron/daily-settle")
async def cron_daily_settle(request: Request):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="unauthorized")
    try:
        body = await request.json()
    except Exception:
        body = {}
    run_id = request.headers.get("X-Webhook-Id") or (body or {}).get("run_id") or secrets.token_hex(8)
    day = (now_utc() - timedelta(days=1)).strftime("%Y-%m-%d")
    asyncio.create_task(run_daily_settlement(run_id, day))
    return {"ok": True, "accepted": True, "run_id": run_id, "date": day}


@api_router.get("/profit-snapshots")
async def list_profit_snapshots(user: dict = Depends(require_roles("admin")), limit: int = Query(90, le=365)):
    rows = await db.profit_snapshots.find().sort("date", -1).to_list(limit)
    return [serialize(r) for r in rows]


@api_router.post("/profit-snapshots/run")
async def run_profit_snapshot(request: Request, date: Optional[str] = None, user: dict = Depends(require_roles("admin"))):
    day = date or (now_utc() - timedelta(days=1)).strftime("%Y-%m-%d")
    if day > now_utc().strftime("%Y-%m-%d"):
        raise HTTPException(status_code=400, detail="不能为未来日期生成快照")
    snap = await run_daily_settlement("manual-" + secrets.token_hex(6), day)
    await audit(user, "settle.snapshot", "snapshot", day, {"order_count": snap.get("order_count")}, request)
    doc = await db.profit_snapshots.find_one({"date": day})
    return serialize(doc) if doc else {"date": day, "order_count": 0}


def _tx_explorer(network: str, txid: str) -> str:
    return f"https://tronscan.org/#/transaction/{txid}"


@api_router.get("/chain/txs")
async def list_chain_txs(user: dict = Depends(require_roles("admin")), address: Optional[str] = None,
                         order_no: Optional[str] = None, limit: int = Query(100, le=500)):
    q: dict = {}
    if address:
        q["to"] = {"$regex": f"^{re.escape(address)}$", "$options": "i"}
    if order_no:
        q["order_no"] = order_no
    txs = await db.chain_txs.find(q).sort("created_at", -1).to_list(limit)
    return [{**serialize(t), "explorer": _tx_explorer(t["network"], t["tx_id"])} for t in txs]


@api_router.post("/chain/sync")
async def chain_sync_now(user: dict = Depends(require_roles("admin"))):
    pairs = await live_address_list()
    scanned, matched, errors = 0, 0, []
    since = int((now_utc() - timedelta(hours=24)).timestamp() * 1000)
    for ch, addr in pairs:
        before = await db.orders.count_documents({"status": "paid", "wallet_address": addr})
        try:
            scanned += await asyncio.wait_for(sync_address(ch, addr, since), timeout=25)
        except asyncio.TimeoutError:
            errors.append(f"{ch['network']} {addr[:8]}…: 超时（地址交易过多），后台轮询将继续处理")
            continue
        except Exception as e:
            errors.append(f"{ch['network']} {addr[:8]}…: {str(e)[:80]}")
            continue
        matched += await db.orders.count_documents({"status": "paid", "wallet_address": addr}) - before
    return {"channels": len({str(ch["_id"]) for ch, _ in pairs}), "addresses": len(pairs),
            "scanned": scanned, "matched": matched, "errors": errors}


@api_router.get("/chain/addresses")
async def list_address_health(user: dict = Depends(require_roles("admin"))):
    pairs = await live_address_list()
    health = {(h["network"], h["address"]): h async for h in db.address_health.find()}
    meta = {m["address"]: m async for m in db.address_meta.find()}
    pipeline = [{"$match": {"status": "pending"}}, {"$group": {"_id": "$wallet_address", "n": {"$sum": 1}}}]
    pending = {r["_id"]: r["n"] async for r in db.orders.aggregate(pipeline)}
    paid_pipeline = [{"$match": {"status": "paid", "live": True}},
                     {"$group": {"_id": "$wallet_address", "n": {"$sum": 1}, "vol": {"$sum": "$amount_usdt"}, "last": {"$max": "$paid_at"}}}]
    paid = {r["_id"]: r async for r in db.orders.aggregate(paid_pipeline)}
    sweep_pipeline = [{"$group": {"_id": "$address", "total": {"$sum": "$amount"}, "last": {"$max": "$created_at"}, "n": {"$sum": 1}}}]
    sweeps = {r["_id"]: r async for r in db.sweeps.aggregate(sweep_pipeline)}
    rows = []
    for ch, addr in pairs:
        h = health.get((ch["network"], addr), {})
        p = paid.get(addr, {})
        m = meta.get(addr, {})
        sw = sweeps.get(addr, {})
        explorer = f"https://tronscan.org/#/address/{addr}"
        bal = h.get("balance")
        rows.append({"address": addr, "network": ch["network"], "channel_id": str(ch["_id"]), "channel_name": ch["name"],
                     "label": m.get("label", ""), "paused": m.get("status") == "paused",
                     "balance": bal, "balance_at": h.get("balance_at"), "balance_error": h.get("balance_error"),
                     "last_sync_at": h.get("last_sync_at"), "sync_error": h.get("sync_error"),
                     "pending": pending.get(addr, 0), "paid_count": p.get("n", 0), "paid_volume": round(p.get("vol", 0), 2),
                     "last_paid_at": p.get("last"), "explorer": explorer,
                     "swept_total": round(sw.get("total", 0), 2), "swept_count": sw.get("n", 0), "last_swept_at": sw.get("last"),
                     "unswept": round(bal, 2) if bal is not None else None,
                     "status": "error" if h.get("sync_error") or h.get("balance_error") else ("ok" if h.get("last_sync_at") or h.get("balance_at") else "idle")})
    total_balance = round(sum(r["balance"] or 0 for r in rows), 2)
    return {"total_balance": total_balance, "total_swept": round(sum(r["swept_total"] for r in rows), 2),
            "by_network": {n: round(sum(r["balance"] or 0 for r in rows if r["network"] == n), 2) for n in NETWORKS},
            "addresses": rows}


class AddressMeta(BaseModel):
    label: Optional[str] = None
    status: Optional[str] = None  # active | paused


@api_router.patch("/chain/addresses/{address}")
async def update_address_meta(address: str, data: AddressMeta, request: Request, user: dict = Depends(require_roles("admin"))):
    upd = {k: v for k, v in data.model_dump().items() if v is not None}
    if "status" in upd and upd["status"] not in ("active", "paused"):
        raise HTTPException(status_code=400, detail="status 需为 active 或 paused")
    if not upd:
        raise HTTPException(status_code=400, detail="无更新内容")
    await db.address_meta.update_one({"address": address}, {"$set": {**upd, "updated_at": iso(now_utc())}}, upsert=True)
    await audit(user, "address.update", "address", address, upd, request)
    m = await db.address_meta.find_one({"address": address})
    return serialize(m)


class SweepCreate(BaseModel):
    address: str
    network: str
    amount: float
    tx_id: Optional[str] = ""
    cold_wallet: Optional[str] = ""
    note: Optional[str] = ""


@api_router.post("/chain/sweeps")
async def create_sweep(data: SweepCreate, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.amount <= 0:
        raise HTTPException(status_code=400, detail="归集金额必须大于 0")
    if data.network not in NETWORKS:
        raise HTTPException(status_code=400, detail="不支持的网络")
    doc = {**data.model_dump(), "operator": user["email"], "created_at": iso(now_utc())}
    res = await db.sweeps.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit(user, "address.sweep", "address", data.address,
                {"amount": data.amount, "network": data.network, "cold_wallet": data.cold_wallet, "tx_id": data.tx_id}, request)
    return serialize(doc)


@api_router.get("/chain/sweeps")
async def list_sweeps(user: dict = Depends(require_roles("admin")), address: Optional[str] = None, limit: int = Query(100, le=500)):
    q = {"address": address} if address else {}
    rows = await db.sweeps.find(q).sort("created_at", -1).to_list(limit)
    return [serialize(r) for r in rows]


class AddressAdd(BaseModel):
    channel_id: str
    addresses: List[str]
    label: Optional[str] = ""


@api_router.post("/chain/addresses")
async def add_addresses(data: AddressAdd, request: Request, user: dict = Depends(require_roles("admin"))):
    ch = await db.channels.find_one({"_id": ObjectId(data.channel_id)})
    if not ch:
        raise HTTPException(status_code=404, detail="通道不存在")
    merged = normalize_addresses(ch["network"], channel_addresses(ch) + data.addresses, None)
    await db.channels.update_one({"_id": ch["_id"]}, {"$set": {"wallet_addresses": merged, "wallet_address": merged[0] if merged else ""}})
    if data.label:
        for a in data.addresses:
            await db.address_meta.update_one({"address": a.strip()}, {"$set": {"label": data.label}}, upsert=True)
    await audit(user, "address.add", "channel", data.channel_id, {"added": data.addresses, "channel": ch["name"]}, request)
    return {"wallet_addresses": merged}


@api_router.delete("/chain/addresses/{channel_id}/{address}")
async def remove_address(channel_id: str, address: str, request: Request, user: dict = Depends(require_roles("admin"))):
    ch = await db.channels.find_one({"_id": ObjectId(channel_id)})
    if not ch:
        raise HTTPException(status_code=404, detail="通道不存在")
    pending = await db.orders.count_documents({"status": "pending", "wallet_address": address})
    if pending:
        raise HTTPException(status_code=400, detail=f"该地址仍有 {pending} 笔待支付订单，请先暂停并等待订单结束")
    remaining = [a for a in channel_addresses(ch) if a != address]
    await db.channels.update_one({"_id": ch["_id"]}, {"$set": {"wallet_addresses": remaining, "wallet_address": remaining[0] if remaining else ""}})
    await audit(user, "address.remove", "channel", channel_id, {"removed": address, "channel": ch["name"]}, request)
    return {"wallet_addresses": remaining}


@api_router.post("/chain/addresses/refresh")
async def refresh_address_health(user: dict = Depends(require_roles("admin"))):
    pairs = await live_address_list()
    ok, errors = 0, []
    for ch, addr in pairs:
        try:
            bal = await fetch_balance(ch["network"], addr)
            await check_balance_alert(ch["network"], addr, ch, bal)
            await check_sweep_alert(ch["network"], addr, ch, bal)
            await set_address_health(ch["network"], addr, ch, balance=bal, balance_at=iso(now_utc()), balance_error=None)
            ok += 1
        except Exception as e:
            await set_address_health(ch["network"], addr, ch, balance_error=str(e)[:200])
            errors.append(f"{ch['network']} {addr[:8]}…: {str(e)[:80]}")
    return {"ok": ok, "errors": errors}


# ---------------------------------------------------------------------------
# Alerts
# ---------------------------------------------------------------------------
class AlertSettings(BaseModel):
    pct: float = 10.0
    min_abs: float = 100.0
    sync_minutes: int = 10
    sweep_threshold: float = 10000.0
    big_payment: float = 1000.0


class NotifySettingsIn(BaseModel):
    telegram_enabled: bool = False
    telegram_bot_token: str = ""
    telegram_chat_id: str = ""
    notify_alerts: bool = True
    notify_big_payment: bool = True
    notify_sweep: bool = True


@api_router.get("/notify/settings")
async def get_notify_settings(user: dict = Depends(require_roles("admin"))):
    cfg = await notify_settings()
    tok = cfg["telegram_bot_token"]
    cfg["telegram_bot_token_masked"] = (tok[:6] + "…" + tok[-4:]) if len(tok) > 12 else ("已设置" if tok else "")
    cfg["telegram_bot_token"] = ""
    return cfg


@api_router.put("/notify/settings")
async def put_notify_settings(data: NotifySettingsIn, request: Request, user: dict = Depends(require_roles("admin"))):
    old = await notify_settings()
    doc = data.model_dump()
    if not doc["telegram_bot_token"]:
        doc["telegram_bot_token"] = old["telegram_bot_token"]
    doc["telegram_chat_id"] = doc["telegram_chat_id"].strip()
    await db.settings.update_one({"_id": "notify"}, {"$set": doc}, upsert=True)
    await audit(user, "settings.notify", "settings", "notify", {k: v for k, v in doc.items() if k != "telegram_bot_token"}, request)
    return {"ok": True}


@api_router.post("/notify/test")
async def test_notify(user: dict = Depends(require_roles("admin"))):
    cfg = await notify_settings()
    if not cfg["telegram_bot_token"] or not cfg["telegram_chat_id"]:
        raise HTTPException(status_code=400, detail="请先填写 Bot Token 与 Chat ID")
    try:
        async with httpx.AsyncClient(timeout=10) as hc:
            r = await hc.post(f"https://api.telegram.org/bot{cfg['telegram_bot_token']}/sendMessage",
                              json={"chat_id": cfg["telegram_chat_id"], "text": f"✅ 四方收款平台 Telegram 推送测试成功\n{iso(now_utc())[:19]}", "parse_mode": "HTML"})
            body = r.json()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"请求 Telegram 失败: {e}")
    if not body.get("ok"):
        raise HTTPException(status_code=400, detail=f"Telegram 返回错误: {body.get('description', body)}")
    return {"ok": True}


@api_router.post("/notify/daily-report")
async def send_daily_report(date: Optional[str] = None, user: dict = Depends(require_roles("admin"))):
    day = date or now_utc().strftime("%Y-%m-%d")
    if day > now_utc().strftime("%Y-%m-%d"):
        raise HTTPException(status_code=400, detail="不能预览未来日期")
    snap = await db.profit_snapshots.find_one({"date": day}) or await build_profit_snapshot(day)
    text = snapshot_report_text(snap)
    cfg = await notify_settings()
    configured = bool(cfg["telegram_bot_token"] and cfg["telegram_chat_id"])
    sent = bool(await send_telegram(text, "report")) if configured else False
    return {"text": text, "telegram_configured": configured, "sent": sent, "date": day}


@api_router.get("/notify/log")
async def notify_log(user: dict = Depends(require_roles("admin")), limit: int = Query(50, le=200)):
    rows = await db.notify_log.find().sort("created_at", -1).to_list(limit)
    return [serialize(r) for r in rows]


@api_router.get("/alerts")
async def list_alerts(user: dict = Depends(require_roles("admin")), unread: bool = False, limit: int = Query(50, le=200)):
    q = {"read": False} if unread else {}
    rows = await db.alerts.find(q).sort("created_at", -1).to_list(limit)
    unread_count = await db.alerts.count_documents({"read": False})
    return {"unread": unread_count, "alerts": [serialize(r) for r in rows]}


@api_router.post("/alerts/read-all")
async def read_all_alerts(user: dict = Depends(require_roles("admin"))):
    await db.alerts.update_many({"read": False}, {"$set": {"read": True}})
    return {"ok": True}


@api_router.post("/alerts/{alert_id}/read")
async def read_alert(alert_id: str, user: dict = Depends(require_roles("admin"))):
    await db.alerts.update_one({"_id": ObjectId(alert_id)}, {"$set": {"read": True}})
    return {"ok": True}


@api_router.get("/alerts/settings")
async def get_alert_settings(user: dict = Depends(require_roles("admin"))):
    return await alert_settings()


@api_router.put("/alerts/settings")
async def put_alert_settings(data: AlertSettings, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.pct <= 0 or data.min_abs < 0 or data.sync_minutes < 1 or data.sweep_threshold < 0 or data.big_payment < 0:
        raise HTTPException(status_code=400, detail="阈值不合法")
    await db.settings.update_one({"_id": "alerts"}, {"$set": data.model_dump()}, upsert=True)
    await audit(user, "settings.alerts", "settings", "alerts", data.model_dump(), request)
    return data.model_dump()


# ---------------------------------------------------------------------------
# Callbacks
# ---------------------------------------------------------------------------
@api_router.get("/callbacks")
async def list_callbacks(user: dict = Depends(get_current_user), status: Optional[str] = None,
                         limit: int = Query(200, le=1000)):
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    if status:
        q["status"] = status
    cbs = await db.callbacks.find(q).sort("created_at", -1).to_list(limit)
    return [serialize(c) for c in cbs]


@api_router.post("/callbacks/{callback_id}/retry")
async def retry_callback(callback_id: str, user: dict = Depends(get_current_user)):
    cb = await db.callbacks.find_one({"_id": ObjectId(callback_id)})
    if not cb:
        raise HTTPException(status_code=404, detail="回调记录不存在")
    scope = await merchant_ids_for_scope(user)
    if scope is not None and ObjectId(cb["merchant_id"]) not in scope:
        raise HTTPException(status_code=403, detail="无权限")
    await db.callbacks.update_one({"_id": cb["_id"]}, {"$set": {"status": "retrying", "next_at": iso(now_utc())}})
    cb = await db.callbacks.find_one({"_id": cb["_id"]})
    await deliver_callback(cb)
    cb = await db.callbacks.find_one({"_id": cb["_id"]})
    return serialize(cb)


@api_router.post("/callbacks/test-receiver")
async def callback_test_receiver(request: Request):
    body = await request.json()
    await db.callback_inbox.insert_one({"payload": body, "received_at": iso(now_utc())})
    return Response(content="success", media_type="text/plain")


@api_router.get("/callbacks/sign-preview")
async def callback_sign_preview(user: dict = Depends(require_roles("merchant"))):
    sample = {"order_no": "SF20260601120000", "amount_usdt": 100.0, "pay_amount": 100.0037,
              "network": "TRC20", "status": "paid", "tx_id": "a" * 64,
              "paid_at": iso(now_utc()), "remark": "", "timestamp": int(now_utc().timestamp())}
    sample["sign"] = sign_payload(sample, user.get("api_secret", ""))
    items = sorted((k, v) for k, v in sample.items() if k != "sign" and v is not None and v != "")
    return {"payload": sample, "sign_string": "&".join(f"{k}={canon(v)}" for k, v in items),
            "algorithm": "HMAC-SHA256(api_secret, sign_string) -> hex"}


# ---------------------------------------------------------------------------
# Audit logs
# ---------------------------------------------------------------------------
@api_router.get("/audit-logs")
async def list_audit_logs(user: dict = Depends(require_roles("admin")), action: Optional[str] = None,
                          actor: Optional[str] = None, actor_id: Optional[str] = None, limit: int = Query(200, le=1000)):
    q: dict = {}
    if action:
        q["action"] = {"$regex": f"^{action}", "$options": "i"}
    if actor:
        q["actor_email"] = {"$regex": actor, "$options": "i"}
    if actor_id:
        q["actor_id"] = actor_id
    logs = await db.audit_logs.find(q).sort("created_at", -1).to_list(limit)
    return [serialize(l) for l in logs]


@api_router.get("/operators")
async def list_operators(user: dict = Depends(require_roles("admin"))):
    ops = await db.users.find({"role": "operator"}).sort("created_at", -1).to_list(200)
    result = []
    for o in ops:
        oid = str(o["_id"])
        count = await db.audit_logs.count_documents({"actor_id": oid})
        last = await db.audit_logs.find({"actor_id": oid}).sort("created_at", -1).limit(1).to_list(1)
        result.append({
            "id": oid, "name": o["name"], "email": o["email"],
            "status": o.get("status", "active"), "permissions": o.get("permissions", []),
            "totp_enabled": bool(o.get("totp_enabled")), "totp_required": bool(o.get("totp_required")),
            "action_count": count, "last_action_at": last[0]["created_at"] if last else None,
        })
    return result


@api_router.get("/auth/lockouts")
async def list_lockouts(user: dict = Depends(require_roles("admin"))):
    recs = await db.login_attempts.find().sort("last_fail_at", -1).to_list(200)
    now = iso(now_utc())
    return [{**serialize(r), "locked": bool(r.get("locked_until") and r["locked_until"] > now)} for r in recs]


@api_router.delete("/auth/lockouts/{identifier:path}")
async def clear_lockout(identifier: str, request: Request, user: dict = Depends(require_roles("admin"))):
    await db.login_attempts.delete_one({"identifier": identifier})
    await audit(user, "auth.unlock", "login", identifier, {}, request)
    return {"ok": True}


# ---------------------------------------------------------------------------
# Channels
# ---------------------------------------------------------------------------
@api_router.get("/channels")
async def list_channels(user: dict = Depends(get_current_user)):
    channels = await db.channels.find().sort("created_at", -1).to_list(200)
    return [await channel_view(c) for c in channels]


def validate_wallet(network: str, address: str):
    if not address:
        return
    if not (address.startswith("T") and len(address) == 34):
        raise HTTPException(status_code=400, detail="TRC20 地址格式不正确（T 开头 34 位）")


def normalize_addresses(network: str, addrs: Optional[List[str]], single: Optional[str]) -> List[str]:
    raw = list(addrs or []) + ([single] if single else [])
    out: List[str] = []
    for a in raw:
        a = (a or "").strip()
        if a and a not in out:
            validate_wallet(network, a)
            out.append(a)
    return out


async def channel_pending_counts(ch: dict) -> dict:
    addrs = channel_addresses(ch)
    if not addrs:
        return {}
    pipeline = [{"$match": {"status": "pending", "wallet_address": {"$in": addrs}}},
                {"$group": {"_id": "$wallet_address", "n": {"$sum": 1}}}]
    return {r["_id"]: r["n"] async for r in db.orders.aggregate(pipeline)}


async def channel_view(ch: dict) -> dict:
    v = serialize(ch)
    v["wallet_addresses"] = channel_addresses(ch)
    v["wallet_address"] = v["wallet_addresses"][0] if v["wallet_addresses"] else ""
    v["pending_by_address"] = await channel_pending_counts(ch)
    return v


@api_router.post("/channels")
async def create_channel(data: ChannelCreate, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.network not in NETWORKS:
        raise HTTPException(status_code=400, detail="不支持的网络")
    addrs = normalize_addresses(data.network, data.wallet_addresses, data.wallet_address)
    doc = data.model_dump()
    doc["wallet_addresses"] = addrs
    doc["wallet_address"] = addrs[0] if addrs else ""
    doc["success_rate"] = round(random.uniform(95, 99.5), 1)
    doc["wallet_balance"] = round(random.uniform(10000, 80000), 2)
    doc["created_at"] = iso(now_utc())
    res = await db.channels.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit(user, "channel.create", "channel", res.inserted_id,
                {"name": doc["name"], "network": doc["network"], "wallet_addresses": addrs}, request)
    return await channel_view(doc)


@api_router.patch("/channels/{channel_id}")
async def update_channel(channel_id: str, data: ChannelUpdate, request: Request,
                         user: dict = Depends(require_roles("admin"))):
    upd = {k: v for k, v in data.model_dump().items() if v is not None}
    if not upd:
        raise HTTPException(status_code=400, detail="无更新内容")
    old = await db.channels.find_one({"_id": ObjectId(channel_id)})
    if not old:
        raise HTTPException(status_code=404, detail="通道不存在")
    if "wallet_addresses" in upd or "wallet_address" in upd:
        addrs = normalize_addresses(old["network"], upd.pop("wallet_addresses", None), upd.pop("wallet_address", None))
        upd["wallet_addresses"] = addrs
        upd["wallet_address"] = addrs[0] if addrs else ""
    await db.channels.update_one({"_id": ObjectId(channel_id)}, {"$set": upd})
    c = await db.channels.find_one({"_id": ObjectId(channel_id)})
    await audit(user, "channel.update", "channel", channel_id,
                {"name": old["name"], "changes": upd, "before": {k: old.get(k) for k in upd}}, request)
    return await channel_view(c)


# ---------------------------------------------------------------------------
# Accounts (merchants & agents)
# ---------------------------------------------------------------------------
@api_router.get("/accounts")
async def list_accounts(user: dict = Depends(get_current_user), role: Optional[str] = None):
    q: dict = {}
    if user["role"] == "admin":
        q["role"] = role if role else {"$in": ["agent", "merchant", "operator"]}
    elif user["role"] == "operator":
        q["role"] = role if role in ("agent", "merchant") else {"$in": ["agent", "merchant"]}
    elif user["role"] == "agent":
        q = {"role": "merchant", "parent_agent_id": str(user["_id"])}
    else:
        raise HTTPException(status_code=403, detail="无权限")
    accounts = await db.users.find(q).sort("created_at", -1).to_list(500)
    return [clean_user(a) for a in accounts]


@api_router.patch("/accounts/me/ip-whitelist")
async def update_my_ip_whitelist(data: IpWhitelist, request: Request, user: dict = Depends(require_roles("merchant"))):
    import ipaddress
    ips: List[str] = []
    for raw in data.ips:
        s = raw.strip()
        if not s:
            continue
        try:
            ipaddress.ip_address(s)
        except ValueError:
            raise HTTPException(status_code=400, detail=f"IP 格式不正确: {s}")
        if s not in ips:
            ips.append(s)
    await db.users.update_one({"_id": user["_id"]}, {"$set": {"ip_whitelist": ips}})
    await audit(user, "merchant.ip_whitelist", "user", user["_id"], {"ips": ips}, request)
    return {"ip_whitelist": ips, "your_ip": client_ip(request)}


@api_router.get("/accounts/me/ip")
async def my_ip(request: Request, user: dict = Depends(get_current_user)):
    return {"ip": client_ip(request)}


@api_router.patch("/accounts/me/notify")
async def update_my_notify(data: NotifySettings, request: Request, user: dict = Depends(require_roles("merchant"))):
    url = (data.notify_url or "").strip()
    if url and not url.startswith(("http://", "https://")):
        raise HTTPException(status_code=400, detail="回调地址必须以 http:// 或 https:// 开头")
    await db.users.update_one({"_id": user["_id"]}, {"$set": {"notify_url": url}})
    await audit(user, "merchant.notify_url", "user", user["_id"], {"notify_url": url}, request)
    return {"notify_url": url}


@api_router.post("/accounts")
async def create_account(data: AccountCreate, request: Request, user: dict = Depends(require_roles("admin", "agent"))):
    if data.role not in ["agent", "merchant", "operator"]:
        raise HTTPException(status_code=400, detail="角色不合法")
    if user["role"] == "agent" and data.role != "merchant":
        raise HTTPException(status_code=403, detail="代理只能创建商户")
    if data.role == "operator":
        bad = [p for p in (data.permissions or []) if p not in OPERATOR_PERMS]
        if bad:
            raise HTTPException(status_code=400, detail=f"未知权限: {bad}")
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="邮箱已存在")

    parent = data.parent_agent_id
    if user["role"] == "agent":
        parent = str(user["_id"])

    doc = {
        "email": email,
        "password_hash": hash_password(data.password),
        "name": data.name,
        "role": data.role,
        "rate": data.rate,
        "parent_agent_id": parent,
        "balance": 0.0,
        "frozen_balance": 0.0,
        "status": "active",
        "created_at": iso(now_utc()),
    }
    if data.role == "merchant":
        doc["api_key"] = gen_api_key()
        doc["api_secret"] = gen_api_secret()
    if data.role == "operator":
        doc["permissions"] = data.permissions or []
        doc["rate"] = 0
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    await audit(user, "account.create", "user", res.inserted_id,
                {"email": email, "role": data.role, "name": data.name, "rate": data.rate, "permissions": doc.get("permissions")}, request)
    return clean_user(doc)


@api_router.get("/accounts/permissions")
async def list_permissions(user: dict = Depends(require_roles("admin"))):
    return OPERATOR_PERMS


@api_router.patch("/accounts/{account_id}/permissions")
async def update_permissions(account_id: str, data: PermissionsUpdate, request: Request, user: dict = Depends(require_roles("admin"))):
    acc = await db.users.find_one({"_id": ObjectId(account_id), "role": "operator"})
    if not acc:
        raise HTTPException(status_code=404, detail="操作员不存在")
    bad = [p for p in data.permissions if p not in OPERATOR_PERMS]
    if bad:
        raise HTTPException(status_code=400, detail=f"未知权限: {bad}")
    await db.users.update_one({"_id": acc["_id"]}, {"$set": {"permissions": data.permissions}})
    await audit(user, "account.permissions", "user", account_id, {"target_email": acc["email"], "permissions": data.permissions}, request)
    return {"permissions": data.permissions}


@api_router.get("/accounts/{account_id}")
async def get_account(account_id: str, user: dict = Depends(get_current_user)):
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    if user["role"] == "agent" and acc.get("parent_agent_id") != str(user["_id"]):
        raise HTTPException(status_code=403, detail="无权限")
    if user["role"] == "merchant" and acc["_id"] != user["_id"]:
        raise HTTPException(status_code=403, detail="无权限")
    return clean_user(acc)


@api_router.post("/accounts/{account_id}/adjust")
async def adjust_balance(account_id: str, data: BalanceAdjust, request: Request,
                         user: dict = Depends(require_roles("admin"))):
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    amt = round(data.amount, 4)
    if data.action == "credit":
        await db.users.update_one({"_id": acc["_id"]}, {"$inc": {"balance": amt}})
    elif data.action == "debit":
        await db.users.update_one({"_id": acc["_id"]}, {"$inc": {"balance": -amt}})
    elif data.action == "freeze":
        await db.users.update_one({"_id": acc["_id"]},
                                  {"$inc": {"balance": -amt, "frozen_balance": amt}})
    elif data.action == "unfreeze":
        await db.users.update_one({"_id": acc["_id"]},
                                  {"$inc": {"balance": amt, "frozen_balance": -amt}})
    else:
        raise HTTPException(status_code=400, detail="非法操作")
    await db.balance_logs.insert_one({
        "account_id": account_id, "action": data.action, "amount": amt,
        "remark": data.remark, "operator": user["email"], "created_at": iso(now_utc()),
    })
    updated = await db.users.find_one({"_id": acc["_id"]})
    _cat = {"credit": "adjust_credit", "debit": "adjust_debit", "freeze": "adjust_freeze", "unfreeze": "adjust_unfreeze"}[data.action]
    _delta = {"credit": amt, "debit": -amt, "freeze": -amt, "unfreeze": amt}[data.action]
    await record_ledger(account_id, _cat, _delta, updated.get("balance", 0), updated.get("frozen_balance", 0),
                        "adjust", None, remark=data.remark or "", operator=user["email"])
    await audit(user, f"balance.{data.action}", "user", account_id,
                {"target_email": acc["email"], "amount": amt, "remark": data.remark,
                 "balance_before": round(acc.get("balance", 0), 4), "balance_after": round(updated.get("balance", 0), 4)}, request)
    return clean_user(updated)


@api_router.post("/accounts/{account_id}/toggle-status")
async def toggle_status(account_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    new_status = "frozen" if acc.get("status") == "active" else "active"
    await db.users.update_one({"_id": acc["_id"]}, {"$set": {"status": new_status}})
    await audit(user, "account.freeze" if new_status == "frozen" else "account.unfreeze", "user", account_id,
                {"target_email": acc["email"]}, request)
    return {"status": new_status}


@api_router.post("/accounts/{account_id}/reset-api-key")
async def reset_api_key(account_id: str, request: Request, user: dict = Depends(require_roles("admin", "agent"))):
    acc = await db.users.find_one({"_id": ObjectId(account_id), "role": "merchant"})
    if not acc:
        raise HTTPException(status_code=404, detail="商户不存在")
    if user["role"] == "agent" and acc.get("parent_agent_id") != str(user["_id"]):
        raise HTTPException(status_code=403, detail="无权限")
    key, sec = gen_api_key(), gen_api_secret()
    await db.users.update_one({"_id": acc["_id"]}, {"$set": {"api_key": key, "api_secret": sec}})
    await audit(user, "account.reset_api_key", "user", account_id, {"target_email": acc["email"]}, request)
    return {"api_key": key, "api_secret": sec}


@api_router.post("/accounts/{account_id}/require-2fa")
async def set_require_2fa(account_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    if acc["role"] not in ("admin", "operator"):
        raise HTTPException(status_code=400, detail="仅管理员/操作员账户支持强制两步验证")
    new_val = not acc.get("totp_required")
    await db.users.update_one({"_id": acc["_id"]}, {"$set": {"totp_required": new_val}})
    await audit(user, "account.require_2fa" if new_val else "account.unrequire_2fa", "user", account_id,
                {"target_email": acc["email"], "required": new_val}, request)
    return {"totp_required": new_val}


# ---------------------------------------------------------------------------
# Settlements / reconciliation
# ---------------------------------------------------------------------------
@api_router.get("/settlements/reconciliation")
async def reconciliation(user: dict = Depends(get_current_user)):
    scope = await merchant_ids_for_scope(user)
    q = {"status": "paid"}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    orders = await db.orders.find(q).to_list(5000)

    by_merchant: dict = {}
    for o in orders:
        m = by_merchant.setdefault(o["merchant_id"], {
            "merchant_id": o["merchant_id"], "merchant_name": o["merchant_name"],
            "order_count": 0, "total_amount": 0.0, "total_fee": 0.0,
            "total_receive": 0.0, "total_commission": 0.0, "total_platform": 0.0,
        })
        m["order_count"] += 1
        m["total_amount"] += o["amount_usdt"]
        m["total_fee"] += o.get("fee", 0)
        m["total_receive"] += o.get("merchant_receive", 0)
        m["total_commission"] += o.get("agent_commission", 0)
        m["total_platform"] += o.get("platform_profit", 0)
    rows = []
    for m in by_merchant.values():
        for k in ["total_amount", "total_fee", "total_receive", "total_commission", "total_platform"]:
            m[k] = round(m[k], 2)
        rows.append(m)
    rows.sort(key=lambda x: x["total_amount"], reverse=True)
    return rows


@api_router.get("/settlements/payouts")
async def list_payouts(user: dict = Depends(get_current_user)):
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    payouts = await db.payouts.find(q).sort("created_at", -1).to_list(500)
    return [serialize(p) for p in payouts]


class PayoutCreate(BaseModel):
    amount: float


@api_router.post("/settlements/payouts")
async def request_payout(data: PayoutCreate, user: dict = Depends(get_current_user)):
    if user["role"] not in ["merchant", "agent"]:
        raise HTTPException(status_code=403, detail="仅商户/代理可提现")
    if data.amount <= 0 or data.amount > user.get("balance", 0):
        raise HTTPException(status_code=400, detail="余额不足")
    doc = {
        "merchant_id": str(user["_id"]),
        "merchant_name": user["name"],
        "amount": round(data.amount, 2),
        "status": "pending",
        "created_at": iso(now_utc()),
    }
    await db.users.update_one({"_id": user["_id"]},
                              {"$inc": {"balance": -data.amount, "frozen_balance": data.amount}})
    res = await db.payouts.insert_one(doc)
    doc["_id"] = res.inserted_id
    _u = await db.users.find_one({"_id": user["_id"]})
    await record_ledger(str(user["_id"]), "payout_freeze", -round(data.amount, 2),
                        _u.get("balance", 0), _u.get("frozen_balance", 0),
                        "payout", res.inserted_id, remark="提现申请，冻结待审核")
    return serialize(doc)


@api_router.post("/settlements/payouts/{payout_id}/approve")
async def approve_payout(payout_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    p = await db.payouts.find_one({"_id": ObjectId(payout_id)})
    if not p or p["status"] != "pending":
        raise HTTPException(status_code=400, detail="提现单不存在或已处理")
    await db.payouts.update_one({"_id": p["_id"]},
                                {"$set": {"status": "approved", "approved_at": iso(now_utc())}})
    await db.users.update_one({"_id": ObjectId(p["merchant_id"])},
                              {"$inc": {"frozen_balance": -p["amount"]}})
    _u = await db.users.find_one({"_id": ObjectId(p["merchant_id"])})
    await record_ledger(p["merchant_id"], "payout_done", 0.0,
                        _u.get("balance", 0), _u.get("frozen_balance", 0),
                        "payout", payout_id, remark=f"提现 {p['amount']} USDT 已放款", operator=user["email"])
    await audit(user, "payout.approve", "payout", payout_id,
                {"merchant_name": p["merchant_name"], "amount": p["amount"]}, request)
    return {"ok": True}


@api_router.post("/settlements/payouts/{payout_id}/reject")
async def reject_payout(payout_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    p = await db.payouts.find_one({"_id": ObjectId(payout_id)})
    if not p or p["status"] != "pending":
        raise HTTPException(status_code=400, detail="提现单不存在或已处理")
    await db.payouts.update_one({"_id": p["_id"]}, {"$set": {"status": "rejected"}})
    await db.users.update_one({"_id": ObjectId(p["merchant_id"])},
                              {"$inc": {"balance": p["amount"], "frozen_balance": -p["amount"]}})
    _u = await db.users.find_one({"_id": ObjectId(p["merchant_id"])})
    await record_ledger(p["merchant_id"], "payout_reject", round(p["amount"], 2),
                        _u.get("balance", 0), _u.get("frozen_balance", 0),
                        "payout", payout_id, remark="提现被驳回，退回余额", operator=user["email"])
    await audit(user, "payout.reject", "payout", payout_id,
                {"merchant_name": p["merchant_name"], "amount": p["amount"]}, request)
    return {"ok": True}


# ---------------------------------------------------------------------------
# Reports (xlsx)
# ---------------------------------------------------------------------------
@api_router.get("/reports/agent-commission")
async def agent_commission_report(month: str, agent_id: Optional[str] = None,
                                  user: dict = Depends(require_roles("admin", "agent"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    import io

    if user["role"] == "agent":
        agent = user
    else:
        if not agent_id:
            raise HTTPException(status_code=400, detail="请选择代理")
        agent = await db.users.find_one({"_id": ObjectId(agent_id), "role": "agent"})
        if not agent:
            raise HTTPException(status_code=404, detail="代理不存在")
    try:
        start = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
    end = (start.replace(year=start.year + 1, month=1) if start.month == 12 else start.replace(month=start.month + 1))

    merchants = await db.users.find({"role": "merchant", "parent_agent_id": str(agent["_id"])}).to_list(500)
    m_ids = [str(m["_id"]) for m in merchants]
    orders = await db.orders.find({"merchant_id": {"$in": m_ids}, "status": "paid",
                                   "paid_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("paid_at", 1).to_list(100000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F766E")

    def write_sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    total = {"n": len(orders), "amt": sum(o["amount_usdt"] for o in orders), "fee": sum(o.get("fee", 0) for o in orders),
             "comm": sum(o.get("agent_commission", 0) for o in orders), "recv": sum(o.get("merchant_receive", 0) for o in orders)}
    ws = wb.active
    ws.title = "月度汇总"
    write_sheet(ws, ["项目", "数值"], [
        ["代理", agent["name"]], ["代理邮箱", agent["email"]], ["月份", month],
        ["代理费率", agent.get("rate", 0)], ["下级商户数", len(merchants)],
        ["成功订单数", total["n"]], ["交易流水 (USDT)", round(total["amt"], 2)],
        ["手续费合计 (USDT)", round(total["fee"], 2)], ["商户实收合计 (USDT)", round(total["recv"], 2)],
        ["代理分润合计 (USDT)", round(total["comm"], 2)],
        ["生成时间", now_utc().strftime("%Y-%m-%d %H:%M UTC")]], [24, 36])

    by_m: dict = {}
    for o in orders:
        d = by_m.setdefault(o["merchant_id"], {"name": o["merchant_name"], "n": 0, "amt": 0.0, "fee": 0.0, "comm": 0.0, "recv": 0.0})
        d["n"] += 1; d["amt"] += o["amount_usdt"]; d["fee"] += o.get("fee", 0); d["comm"] += o.get("agent_commission", 0); d["recv"] += o.get("merchant_receive", 0)
    m_rate = {str(m["_id"]): m.get("rate", 0) for m in merchants}
    write_sheet(wb.create_sheet("按商户明细"), ["商户", "商户费率", "订单数", "交易流水", "手续费", "商户实收", "代理分润"],
                [[d["name"], m_rate.get(mid, ""), d["n"], round(d["amt"], 2), round(d["fee"], 2), round(d["recv"], 2), round(d["comm"], 2)]
                 for mid, d in sorted(by_m.items(), key=lambda x: -x[1]["amt"])], [22, 12, 10, 16, 16, 16, 16])

    write_sheet(wb.create_sheet("逐单明细"), ["订单号", "商户", "网络", "金额", "手续费", "商户实收", "代理分润", "TxID", "支付时间"],
                [[o["order_no"], o["merchant_name"], o["network"], o["amount_usdt"], o.get("fee", 0), o.get("merchant_receive", 0),
                  o.get("agent_commission", 0), o.get("tx_id") or "", (o.get("paid_at") or "")[:19].replace("T", " ")] for o in orders],
                [26, 18, 10, 12, 12, 12, 12, 40, 20])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"commission_{agent['name']}_{month}.xlsx"
    from urllib.parse import quote
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@api_router.get("/reports/merchant-statement")
async def merchant_statement(month: str, merchant_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    import io

    if user["role"] == "merchant":
        merchant = user
    else:
        if user["role"] == "operator" and "reports.export" not in (user.get("permissions") or []):
            raise HTTPException(status_code=403, detail="无权限访问")
        if not merchant_id:
            raise HTTPException(status_code=400, detail="请选择商户")
        merchant = await db.users.find_one({"_id": ObjectId(merchant_id), "role": "merchant"})
        if not merchant:
            raise HTTPException(status_code=404, detail="商户不存在")
        if user["role"] == "agent" and merchant.get("parent_agent_id") != str(user["_id"]):
            raise HTTPException(status_code=403, detail="无权限")
    try:
        start = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
    end = (start.replace(year=start.year + 1, month=1) if start.month == 12 else start.replace(month=start.month + 1))
    mid = str(merchant["_id"])
    orders = await db.orders.find({"merchant_id": mid, "status": "paid", "paid_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("paid_at", 1).to_list(100000)
    payouts = await db.payouts.find({"merchant_id": mid, "created_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("created_at", 1).to_list(10000)
    adjustments = await db.audit_logs.find({"action": {"$regex": "^balance\\."}, "target_id": mid,
                                            "created_at": {"$gte": iso(start), "$lt": iso(end)}}).sort("created_at", 1).to_list(10000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0E7490")

    def write_sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [18] * len(headers), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    amt = sum(o["amount_usdt"] for o in orders); fee = sum(o.get("fee", 0) for o in orders); recv = sum(o.get("merchant_receive", 0) for o in orders)
    approved = sum(p["amount"] for p in payouts if p["status"] == "approved")
    by_net: dict = {}
    for o in orders:
        d = by_net.setdefault(o["network"], {"n": 0, "amt": 0.0, "fee": 0.0})
        d["n"] += 1; d["amt"] += o["amount_usdt"]; d["fee"] += o.get("fee", 0)
    ws = wb.active; ws.title = "月度汇总"
    write_sheet(ws, ["项目", "数值"], [
        ["商户", merchant["name"]], ["商户邮箱", merchant["email"]], ["月份", month], ["商户费率", merchant.get("rate", 0)],
        ["成功订单数", len(orders)], ["收款流水 (USDT)", round(amt, 2)], ["手续费合计 (USDT)", round(fee, 2)], ["实收合计 (USDT)", round(recv, 2)],
        ["提现申请笔数", len(payouts)], ["已通过提现 (USDT)", round(approved, 2)],
        ["当前可用余额 (USDT)", round(merchant.get("balance", 0), 2)], ["当前冻结余额 (USDT)", round(merchant.get("frozen_balance", 0), 2)],
        ["生成时间", now_utc().strftime("%Y-%m-%d %H:%M UTC")]] + [[f"{n} 订单数/流水/手续费", f"{d['n']} / {round(d['amt'], 2)} / {round(d['fee'], 2)}"] for n, d in by_net.items()], [26, 40])
    write_sheet(wb.create_sheet("收款明细"), ["订单号", "商户单号", "网络", "订单金额", "实付金额", "手续费", "实收", "TxID", "收款地址", "支付时间"],
                [[o["order_no"], o.get("out_trade_no") or "", o["network"], o["amount_usdt"], o.get("pay_amount", o["amount_usdt"]), o.get("fee", 0), o.get("merchant_receive", 0),
                  o.get("tx_id") or "", o["wallet_address"], (o.get("paid_at") or "")[:19].replace("T", " ")] for o in orders], [26, 18, 8, 12, 12, 10, 12, 40, 36, 20])
    write_sheet(wb.create_sheet("提现记录"), ["申请时间", "金额", "网络", "提现地址", "状态", "审核时间"],
                [[(p.get("created_at") or "")[:19].replace("T", " "), p["amount"], p.get("network", ""), p.get("address", ""), p["status"], (p.get("approved_at") or "")[:19].replace("T", " ")] for p in payouts], [20, 12, 8, 40, 10, 20])
    write_sheet(wb.create_sheet("余额调整"), ["时间", "操作", "金额", "余额前", "余额后", "操作人", "备注"],
                [[(a["created_at"] or "")[:19].replace("T", " "), a["action"], a["detail"].get("amount"), a["detail"].get("balance_before"), a["detail"].get("balance_after"), a["actor_email"], a["detail"].get("remark", "")] for a in adjustments], [20, 16, 12, 12, 12, 24, 24])

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"statement_{merchant['name']}_{month}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------------------------------------------------------------------------
# Balance ledger (merchant / agent running statement)
# ---------------------------------------------------------------------------
async def resolve_ledger_account(user: dict, account_id: Optional[str]) -> ObjectId:
    if user["role"] == "merchant":
        return user["_id"]
    if user["role"] == "agent":
        if not account_id:
            return user["_id"]
        acc = await db.users.find_one({"_id": ObjectId(account_id)})
        if not acc or (acc["_id"] != user["_id"] and acc.get("parent_agent_id") != str(user["_id"])):
            raise HTTPException(status_code=403, detail="无权限查看该账户")
        return acc["_id"]
    # admin / operator
    if not account_id:
        raise HTTPException(status_code=400, detail="请选择账户")
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    return acc["_id"]


@api_router.get("/balance/ledger")
async def balance_ledger(account_id: Optional[str] = None, category: Optional[str] = None,
                         month: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                         limit: int = Query(300, le=2000),
                         user: dict = Depends(get_current_user)):
    acc_id = await resolve_ledger_account(user, account_id)
    q: dict = {"account_id": str(acc_id)}
    if category:
        q["category"] = category
    if start and end:
        try:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    elif month:
        try:
            ms = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
        me = (ms.replace(year=ms.year + 1, month=1) if ms.month == 12 else ms.replace(month=ms.month + 1))
        q["created_at"] = {"$gte": iso(ms), "$lt": iso(me)}
    rows = await db.balance_ledger.find(q).sort("created_at", -1).to_list(limit)
    acc = await db.users.find_one({"_id": acc_id})
    return {
        "account": {"id": str(acc_id), "name": acc.get("name"), "email": acc.get("email"),
                    "role": acc.get("role"), "balance": acc.get("balance", 0),
                    "frozen_balance": acc.get("frozen_balance", 0)},
        "entries": [serialize(r) for r in rows],
    }


LEDGER_CAT_LABELS = {
    "order_credit": "收款入账", "commission": "代理分润", "payout_freeze": "提现冻结",
    "payout_done": "提现完成", "payout_reject": "提现驳回退回", "adjust_credit": "人工加款",
    "adjust_debit": "人工扣款", "adjust_freeze": "冻结", "adjust_unfreeze": "解冻",
}


@api_router.get("/reports/balance-ledger")
async def balance_ledger_report(account_id: Optional[str] = None, month: Optional[str] = None,
                                start: Optional[str] = None, end: Optional[str] = None,
                                user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    import io

    acc_id = await resolve_ledger_account(user, account_id)
    acc = await db.users.find_one({"_id": acc_id})
    q: dict = {"account_id": str(acc_id)}
    title_suffix = "全部"
    if start and end:
        try:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
        title_suffix = f"{start}_{end}"
    elif month:
        try:
            ms = datetime.strptime(month, "%Y-%m").replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(status_code=400, detail="month 格式应为 YYYY-MM")
        me = (ms.replace(year=ms.year + 1, month=1) if ms.month == 12 else ms.replace(month=ms.month + 1))
        q["created_at"] = {"$gte": iso(ms), "$lt": iso(me)}
        title_suffix = month
    rows = await db.balance_ledger.find(q).sort("created_at", 1).to_list(100000)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws = wb.active; ws.title = "资金流水"
    ws.append(["时间", "类型", "变动金额(USDT)", "变动后余额", "变动后冻结", "关联", "备注", "操作人"])
    for c in ws[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for r in rows:
        ws.append([(r.get("created_at") or "")[:19].replace("T", " "), LEDGER_CAT_LABELS.get(r["category"], r["category"]),
                   r.get("amount"), r.get("balance_after"), r.get("frozen_after"),
                   f"{r.get('ref_type') or ''} {(r.get('ref_id') or '')}".strip(), r.get("remark", ""), r.get("operator") or ""])
    for i, w in enumerate([20, 14, 16, 14, 14, 30, 34, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"ledger_{acc.get('name', '')}_{title_suffix}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@api_router.get("/reports/profit-snapshot")
async def profit_snapshot_report(date: str, user: dict = Depends(require_roles("admin"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    import io
    snap = await db.profit_snapshots.find_one({"date": date}) or await build_profit_snapshot(date)
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    ws = wb.active; ws.title = "日度汇总"
    ws.append([f"分润对账日期：{date}"])
    ws.append(["订单数", "收款总额", "手续费", "商户实收", "代理分润", "平台利润", "补偿结算"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    ws.append([snap["order_count"], snap["total_amount"], snap["total_fee"], snap["total_merchant_receive"],
               snap["total_agent_commission"], snap["total_platform_profit"], snap.get("backfilled", 0)])
    for i, w in enumerate([12, 16, 14, 16, 14, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws2 = wb.create_sheet("按商户明细")
    ws2.append(["商户", "订单数", "收款金额", "手续费", "商户实收", "代理分润", "平台利润"])
    for c in ws2[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    for m in snap.get("merchants", []):
        ws2.append([m["merchant_name"], m["orders"], m["amount"], m["fee"], m["receive"], m["commission"], m["platform"]])
    for i, w in enumerate([22, 10, 16, 14, 16, 14, 14], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"; ws2.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"profit_snapshot_{date}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


ORDER_STATUS_LABELS = {"pending": "待支付", "confirming": "确认中", "paid": "已支付", "expired": "超时未付", "rejected": "已驳回"}


async def generate_orders_xlsx(user, status=None, network=None, keyword=None, merchant_id=None,
                               date=None, start=None, end=None, ids=None, progress=None):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io
    scope = await merchant_ids_for_scope(user)
    q: dict = {}
    if scope is not None:
        q["merchant_id"] = {"$in": [str(i) for i in scope]}
    id_list = [x for x in (ids.split(",") if ids else []) if x]
    if id_list:
        try:
            q["_id"] = {"$in": [ObjectId(x) for x in id_list]}
        except Exception:
            raise HTTPException(status_code=400, detail="订单 ID 格式不正确")
    else:
        if merchant_id and (scope is None or ObjectId(merchant_id) in scope):
            q["merchant_id"] = str(merchant_id)
        if start and end:
            ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
            q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
        elif date:
            ds = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
            q["created_at"] = {"$gte": iso(ds), "$lt": iso(ds + timedelta(days=1))}
        if status:
            q["status"] = status
        if network:
            q["network"] = network
        if keyword:
            kw = keyword.strip()
            q["$or"] = [{"order_no": {"$regex": kw, "$options": "i"}}, {"tx_id": {"$regex": kw, "$options": "i"}},
                        {"out_trade_no": {"$regex": kw, "$options": "i"}}, {"wallet_address": {"$regex": kw, "$options": "i"}},
                        {"merchant_name": {"$regex": kw, "$options": "i"}}]
            if re.fullmatch(r"\d+(\.\d+)?", kw):
                q["$or"] += [{"amount_usdt": float(kw)}, {"pay_amount": float(kw)}]
    filters = []
    if id_list:
        filters.append(f"所选订单 {len(id_list)} 笔")
    else:
        if merchant_id:
            filters.append(f"商户ID={merchant_id}")
        if start and end:
            filters.append(f"区间 {start}~{end}")
        elif date:
            filters.append(f"日期 {date}")
        if status:
            filters.append(f"状态={ORDER_STATUS_LABELS.get(status, status)}")
        if network:
            filters.append(f"网络={network}")
        if keyword:
            filters.append(f"关键词={keyword}")
    summary = "、".join(filters) or "全部订单"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("订单明细")
    for i, w in enumerate([22, 16, 18, 8, 12, 12, 10, 36, 40, 10, 12, 10, 20, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0E7490")
    title = WriteOnlyCell(ws, value=f"订单导出 · 筛选条件：{summary}"); title.font = Font(bold=True, size=12)
    ws.append([title])
    ws.append([f"导出时间：{iso(now_utc())[:19].replace('T', ' ')} (UTC)    操作人：{user.get('email', '')}"])
    ws.append([])
    header = ["订单号", "商户", "商户单号", "网络", "金额(USDT)", "到账金额", "状态", "收款地址", "TxID",
              "手续费", "实收", "回调状态", "创建时间", "支付时间"]
    hrow = []
    for h in header:
        c = WriteOnlyCell(ws, value=h); c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
        hrow.append(c)
    ws.append(hrow)
    total = await db.orders.count_documents(q)
    if progress:
        await progress(0, total)
    count = 0
    async for o in db.orders.find(q).sort("created_at", -1):
        ws.append([o.get("order_no"), o.get("merchant_name"), o.get("out_trade_no") or "", o.get("network"),
                   o.get("amount_usdt"), o.get("pay_amount"), ORDER_STATUS_LABELS.get(o.get("status"), o.get("status")),
                   o.get("wallet_address") or "", o.get("tx_id") or "", o.get("fee"), o.get("merchant_receive"),
                   o.get("callback_status") or "", (o.get("created_at") or "")[:19].replace("T", " "),
                   (o.get("paid_at") or "")[:19].replace("T", " ")])
        count += 1
        if progress and count % 1000 == 0:
            await progress(count, total)
    ws.freeze_panes = "A5"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    tag = "selected" if id_list else (date or (f"{start}_{end}" if start and end else "all"))
    return buf, count, f"orders_{tag}.xlsx"


@api_router.get("/reports/orders")
async def orders_report(user: dict = Depends(get_current_user),
                        status: Optional[str] = None, network: Optional[str] = None,
                        keyword: Optional[str] = None, merchant_id: Optional[str] = None,
                        date: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None,
                        ids: Optional[str] = None):
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    buf, _count, fname = await generate_orders_xlsx(user, status, network, keyword, merchant_id, date, start, end, ids)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


EXPORT_KINDS = {"orders", "snapshots", "ledger"}
LEDGER_CAT = {
    "order_credit": "收款入账", "commission": "代理分润", "payout_freeze": "提现冻结", "payout_done": "提现完成",
    "payout_reject": "提现驳回退回", "adjust_credit": "人工加款", "adjust_debit": "人工扣款",
    "adjust_freeze": "冻结", "adjust_unfreeze": "解冻",
}


async def generate_snapshots_xlsx(start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io
    q = {}
    if start and end:
        q["date"] = {"$gte": start, "$lte": end}
    rows = await db.profit_snapshots.find(q).sort("date", -1).to_list(500)
    wb = Workbook(); ws = wb.active; ws.title = "分润快照"
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="0E7490")
    ws.append(["日期", "订单数", "收款总额", "手续费", "商户实收", "代理分润", "平台利润", "补偿单", "环比%"])
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, hfill, Alignment(horizontal="center")
    for s in rows:
        ws.append([s.get("date"), s.get("order_count"), s.get("total_amount"), s.get("total_fee"),
                   s.get("total_merchant_receive"), s.get("total_agent_commission"), s.get("total_platform_profit"),
                   s.get("backfilled", 0), s.get("platform_profit_change_pct")])
    for i, w in enumerate([12, 10, 14, 12, 14, 12, 12, 8, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, len(rows), f"profit_snapshots_{start or 'all'}.xlsx"


async def generate_ledger_xlsx(user, account_id=None, month=None, start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io
    acc = await db.users.find_one({"_id": ObjectId(account_id)})
    if not acc:
        raise HTTPException(status_code=404, detail="账户不存在")
    q = {"account_id": account_id}
    if start and end:
        ds = datetime.strptime(start, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        de = datetime.strptime(end, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    elif month:
        ds = datetime.strptime(month + "-01", "%Y-%m-%d").replace(tzinfo=timezone.utc)
        de = (ds + timedelta(days=32)).replace(day=1)
        q["created_at"] = {"$gte": iso(ds), "$lt": iso(de)}
    rows = await db.balance_logs.find(q).sort("created_at", -1).to_list(100000)
    wb = Workbook(); ws = wb.active; ws.title = "资金流水"
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="0E7490")
    ws.append([f"账户：{acc.get('name', '')} ({acc.get('email', '')})"])
    ws.append(["时间", "类型", "金额(USDT)", "变动后余额", "关联单号/ID", "备注"])
    for c in ws[2]:
        c.font, c.fill, c.alignment = hf, hfill, Alignment(horizontal="center")
    for r in rows:
        ws.append([(r.get("created_at") or "")[:19].replace("T", " "), LEDGER_CAT.get(r.get("category"), r.get("category")),
                   r.get("amount"), r.get("balance_after"), r.get("ref_id") or "", r.get("remark") or ""])
    for i, w in enumerate([20, 12, 14, 14, 24, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, len(rows), f"ledger_{acc.get('name', '')}_{month or start or 'all'}.xlsx"


# ---------------------------------------------------------------------------
# Background export jobs (large exports -> async job + download link)
# ---------------------------------------------------------------------------
def _require_export(user: dict):
    if user["role"] == "operator" and "reports.export" not in (user.get("permissions") or []):
        raise HTTPException(status_code=403, detail="无报表导出权限")


async def _run_export(job_id: str, user: dict, kind: str, params: dict):
    async def progress(done, total):
        await db.export_jobs.update_one({"_id": ObjectId(job_id)}, {"$set": {"processed": done, "expected": total}})
    try:
        if kind == "orders":
            buf, count, fname = await generate_orders_xlsx(user, progress=progress, **params)
        elif kind == "snapshots":
            buf, count, fname = await generate_snapshots_xlsx(**params)
        else:
            buf, count, fname = await generate_ledger_xlsx(user, **params)
        data = buf.getvalue()
        file_id = await fs.upload_from_stream(f"{job_id}.xlsx", data)
        await db.export_jobs.update_one({"_id": ObjectId(job_id)},
            {"$set": {"status": "done", "row_count": count, "filename": fname, "size_bytes": len(data),
                      "file_id": str(file_id), "processed": count, "expected": count,
                      "finished_at": iso(now_utc()), "expires_at": iso(now_utc() + timedelta(days=7))}})
        if count >= 500:
            kind_label = {"orders": "订单", "snapshots": "分润快照", "ledger": "资金流水"}.get(kind, kind)
            await send_telegram(f"📁 <b>导出完成</b>\n类型：{kind_label}\n行数：{count}\n大小：{round(len(data) / 1024)}KB\n操作人：{user.get('email', '')}", "report")
    except Exception as e:
        await db.export_jobs.update_one({"_id": ObjectId(job_id)},
            {"$set": {"status": "failed", "error": str(e)[:300], "finished_at": iso(now_utc())}})


@api_router.post("/export/{kind}")
async def create_export(kind: str, request: Request, user: dict = Depends(get_current_user)):
    _require_export(user)
    if kind not in EXPORT_KINDS:
        raise HTTPException(status_code=400, detail="不支持的导出类型")
    try:
        params = await request.json()
    except Exception:
        params = {}
    fields = {"orders": ("status", "network", "keyword", "merchant_id", "date", "start", "end", "ids"),
              "snapshots": ("start", "end"),
              "ledger": ("account_id", "month", "start", "end")}[kind]
    allowed = {k: params.get(k) for k in fields if params.get(k)}
    if kind == "ledger" and not allowed.get("account_id"):
        raise HTTPException(status_code=400, detail="导出资金流水需指定 account_id")
    doc = {"kind": kind, "status": "pending", "params": allowed,
           "created_by": str(user["_id"]), "created_email": user["email"], "created_at": iso(now_utc())}
    res = await db.export_jobs.insert_one(doc)
    jid = str(res.inserted_id)
    asyncio.create_task(_run_export(jid, user, kind, allowed))
    await audit(user, "export.create", "export", jid, {"kind": kind, "params": allowed}, request)
    return {"id": jid, "status": "pending"}


@api_router.get("/export/jobs")
async def list_export_jobs(user: dict = Depends(get_current_user)):
    _require_export(user)
    q = {} if user["role"] == "admin" else {"created_by": str(user["_id"])}
    rows = await db.export_jobs.find(q).sort("created_at", -1).to_list(50)
    return [serialize(r) for r in rows]


@api_router.get("/export/jobs/{job_id}/download")
async def download_export_job(job_id: str, user: dict = Depends(get_current_user)):
    from fastapi.responses import StreamingResponse
    from urllib.parse import quote
    import io
    _require_export(user)
    job = await db.export_jobs.find_one({"_id": ObjectId(job_id)})
    if not job or job.get("status") != "done" or not job.get("file_id"):
        raise HTTPException(status_code=404, detail="文件未就绪")
    if user["role"] != "admin" and job.get("created_by") != str(user["_id"]):
        raise HTTPException(status_code=403, detail="无权下载")
    try:
        stream = await fs.open_download_stream(ObjectId(job["file_id"]))
        data = await stream.read()
    except Exception:
        raise HTTPException(status_code=410, detail="导出文件已过期，请重新导出")
    fname = job.get("filename", "orders.xlsx")
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@api_router.get("/profit-snapshots/trend")
async def profit_trend(days: int = Query(30, ge=7, le=90), merchant_id: Optional[str] = None,
                       user: dict = Depends(require_roles("admin"))):
    today = now_utc()
    start_day = (today - timedelta(days=days - 1)).strftime("%Y-%m-%d")
    q: dict = {"status": "paid", "paid_at": {"$gte": start_day + "T00:00:00"}}
    if merchant_id:
        q["merchant_id"] = merchant_id
    buckets: dict = {}
    async for o in db.orders.find(q):
        d = (o.get("paid_at") or "")[:10]
        if not d:
            continue
        b = buckets.setdefault(d, {"platform": 0.0, "commission": 0.0, "amount": 0.0, "orders": 0})
        b["platform"] += o.get("platform_profit", 0); b["commission"] += o.get("agent_commission", 0)
        b["amount"] += o.get("amount_usdt", 0); b["orders"] += 1
    out = []
    for i in range(days):
        d = (today - timedelta(days=days - 1 - i)).strftime("%Y-%m-%d")
        b = buckets.get(d, {"platform": 0, "commission": 0, "amount": 0, "orders": 0})
        out.append({"date": d, "platform_profit": round(b["platform"], 2),
                    "agent_commission": round(b["commission"], 2), "amount": round(b["amount"], 2), "orders": b["orders"]})
    return out


# ---------------------------------------------------------------------------
# Seed
# ---------------------------------------------------------------------------
async def seed():
    await db.users.create_index("email", unique=True)
    await db.chain_txs.create_index("tx_id", unique=True)
    await db.audit_logs.create_index([("created_at", -1)])
    await db.callbacks.create_index([("status", 1), ("next_at", 1)])
    await db.login_attempts.create_index("identifier")
    await db.api_nonces.create_index("created_at", expireAfterSeconds=600)
    await db.api_nonces.create_index([("api_key", 1), ("nonce", 1)])
    await db.orders.create_index([("status", 1), ("wallet_address", 1), ("pay_amount", 1)])
    await db.orders.create_index([("merchant_id", 1), ("out_trade_no", 1)])
    await db.alerts.create_index([("read", 1), ("created_at", -1)])
    await db.balance_ledger.create_index([("account_id", 1), ("created_at", -1)])
    await db.audit_logs.create_index([("actor_id", 1), ("created_at", -1)])
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    admin = await db.users.find_one({"email": admin_email})
    if admin is None:
        await db.users.insert_one({
            "email": admin_email, "password_hash": hash_password(admin_password),
            "name": "平台管理员", "role": "admin", "balance": 0.0,
            "frozen_balance": 0.0, "status": "active", "created_at": iso(now_utc()),
        })
    elif not verify_password(admin_password, admin["password_hash"]):
        await db.users.update_one({"email": admin_email},
                                  {"$set": {"password_hash": hash_password(admin_password)}})

    if await db.channels.count_documents({}) == 0:
        base = [
            {"name": "TRC20 极速通道", "network": "TRC20", "fee_rate": 0.02, "weight": 50,
             "min_amount": 10, "max_amount": 50000, "status": "active"},
            {"name": "TRC20 大额通道", "network": "TRC20", "fee_rate": 0.015, "weight": 30,
             "min_amount": 50, "max_amount": 100000, "status": "active"},
        ]
        for c in base:
            c["success_rate"] = round(random.uniform(96, 99.5), 1)
            c["wallet_balance"] = round(random.uniform(20000, 90000), 2)
            c["created_at"] = iso(now_utc())
        await db.channels.insert_many(base)

    # demo agent + merchants
    if await db.users.count_documents({"role": "agent"}) == 0:
        agent_res = await db.users.insert_one({
            "email": "agent@demo.com", "password_hash": hash_password("agent888"),
            "name": "华东一级代理", "role": "agent", "rate": 0.015,
            "parent_agent_id": None, "balance": 0.0, "frozen_balance": 0.0,
            "status": "active", "created_at": iso(now_utc()),
        })
        agent_id = str(agent_res.inserted_id)
        merchants = [
            {"email": "merchant@demo.com", "name": "星辰科技", "rate": 0.03, "parent": agent_id},
            {"email": "shop2@demo.com", "name": "云端商城", "rate": 0.028, "parent": agent_id},
            {"email": "shop3@demo.com", "name": "极光支付", "rate": 0.035, "parent": None},
        ]
        m_ids = []
        for m in merchants:
            r = await db.users.insert_one({
                "email": m["email"], "password_hash": hash_password("merchant888"),
                "name": m["name"], "role": "merchant", "rate": m["rate"],
                "parent_agent_id": m["parent"], "balance": 0.0, "frozen_balance": 0.0,
                "status": "active", "api_key": gen_api_key(), "api_secret": gen_api_secret(),
                "created_at": iso(now_utc()),
            })
            m_ids.append((str(r.inserted_id), m["name"]))

        # demo orders across last 7 days
        channels = await db.channels.find({"status": "active"}).to_list(10)
        for _ in range(120):
            mid, mname = random.choice(m_ids)
            ch = random.choice(channels)
            amount = round(random.uniform(50, 3000), 2)
            created = now_utc() - timedelta(days=random.randint(0, 6),
                                            hours=random.randint(0, 23),
                                            minutes=random.randint(0, 59))
            paid = random.random() < 0.8
            order = {
                "order_no": gen_order_no(), "merchant_id": mid, "merchant_name": mname,
                "amount_usdt": amount, "network": ch["network"],
                "wallet_address": gen_wallet_address(ch["network"]), "tx_id": None,
                "status": "pending", "channel_id": str(ch["_id"]), "channel_name": ch["name"],
                "remark": "", "created_at": iso(created), "paid_at": None,
                "expires_at": iso(created + timedelta(minutes=15)),
            }
            res = await db.orders.insert_one(order)
            order["_id"] = res.inserted_id
            if paid:
                await mark_paid(order, notify=False)
            elif random.random() < 0.5:
                await db.orders.update_one({"_id": order["_id"]}, {"$set": {"status": "expired"}})
    # one-time backfill of balance ledger from existing paid orders
    if await db.balance_ledger.count_documents({}) == 0 and await db.orders.count_documents({"status": "paid"}) > 0:
        merchants = await db.users.find({"role": "merchant"}).to_list(1000)
        m_agent = {str(m["_id"]): m.get("parent_agent_id") for m in merchants}
        paid = await db.orders.find({"status": "paid"}).sort("paid_at", 1).to_list(100000)
        running: dict = {}
        for o in paid:
            ts = o.get("paid_at") or o.get("created_at") or iso(now_utc())
            mid = o["merchant_id"]
            net = round(o.get("merchant_receive", 0), 4)
            running[mid] = round(running.get(mid, 0.0) + net, 4)
            await db.balance_ledger.insert_one({
                "account_id": str(mid), "category": "order_credit", "amount": net,
                "balance_after": running[mid], "frozen_after": 0.0, "ref_type": "order",
                "ref_id": o.get("order_no"), "remark": f"{o.get('network', '')} 收款 {o.get('amount_usdt')} USDT，手续费 {o.get('fee', 0)}",
                "operator": None, "created_at": ts,
            })
            comm = round(o.get("agent_commission", 0), 4)
            aid = m_agent.get(str(mid))
            if comm > 0 and aid:
                running[aid] = round(running.get(aid, 0.0) + comm, 4)
                await db.balance_ledger.insert_one({
                    "account_id": str(aid), "category": "commission", "amount": comm,
                    "balance_after": running[aid], "frozen_after": 0.0, "ref_type": "order",
                    "ref_id": o.get("order_no"), "remark": f"订单 {o.get('order_no')} 代理分润",
                    "operator": None, "created_at": ts,
                })

    logger.info("Seed complete.")


@app.on_event("startup")
async def on_startup():
    await seed()
    asyncio.create_task(chain_listener())
    asyncio.create_task(callback_worker())
    asyncio.create_task(balance_refresher())


@app.on_event("shutdown")
async def on_shutdown():
    client.close()


app.include_router(api_router)

# Extracted domain routers (late import avoids circular dependency with server defs)
from routes.dev import router as dev_router  # noqa: E402
app.include_router(dev_router)
from routes.auth import router as auth_router  # noqa: E402
app.include_router(auth_router)


RATE_LIMIT_MAX = int(os.environ.get("RATE_LIMIT_MAX", "300"))
RATE_LIMIT_WINDOW = 60
_rl_bucket: dict = {}


@app.middleware("http")
async def security_middleware(request: Request, call_next):
    ip = (request.headers.get("x-forwarded-for", "").split(",")[0].strip()
          or (request.client.host if request.client else "")) or "unknown"
    path = request.url.path
    if path.startswith("/api") and request.method != "OPTIONS":
        now = now_utc().timestamp()
        bucket = _rl_bucket.get(ip)
        if not bucket or now - bucket["start"] >= RATE_LIMIT_WINDOW:
            _rl_bucket[ip] = {"start": now, "count": 1}
        else:
            bucket["count"] += 1
            if bucket["count"] > RATE_LIMIT_MAX:
                return JSONResponse(status_code=429, content={"detail": "请求过于频繁，请稍后再试"})
        if len(_rl_bucket) > 5000:
            cutoff = now - RATE_LIMIT_WINDOW
            for k in [k for k, v in _rl_bucket.items() if v["start"] < cutoff]:
                _rl_bucket.pop(k, None)
    resp = await call_next(request)
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "no-referrer"
    resp.headers["X-XSS-Protection"] = "1; mode=block"
    resp.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    resp.headers["Permissions-Policy"] = "geolocation=(), microphone=(), camera=()"
    return resp


app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000")],
    allow_methods=["*"],
    allow_headers=["*"],
)
