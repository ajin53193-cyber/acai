"""Iteration 15 — profit/orders Excel reports, Telegram daily-report, refactor regression."""
import os
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/")
ADMIN = {"email": "ajin53193@gmail.com", "password": "admin888"}
OPERATOR = {"email": "operator@demo.com", "password": "operator888"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _login(email, password):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=15)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text}"
    return r.json()["token"]


@pytest.fixture(scope="module")
def admin_token():
    return _login(**ADMIN)


@pytest.fixture(scope="module")
def operator_token():
    return _login(**OPERATOR)


@pytest.fixture(scope="module")
def admin_h(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="module")
def op_h(operator_token):
    return {"Authorization": f"Bearer {operator_token}"}


# ---------- Refactor regression: auth/me, order create+pay settles ----------
class TestRefactorRegression:
    def test_auth_me(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/auth/me", headers=admin_h, timeout=10)
        assert r.status_code == 200
        d = r.json()
        assert d["email"] == ADMIN["email"] and d["role"] == "admin"
        assert "id" in d and "password_hash" not in d

    def test_order_create_and_settle(self, admin_h):
        # pick a merchant
        rs = requests.get(f"{BASE_URL}/api/accounts?role=merchant", headers=admin_h, timeout=10)
        assert rs.status_code == 200
        merchants = rs.json()
        assert merchants, "need at least one merchant"
        m = merchants[0]
        before = float(m.get("balance") or 0)
        # create order
        payload = {"amount_usdt": 12.34, "network": "TRC20", "merchant_id": m["id"], "remark": "TEST_it15"}
        r = requests.post(f"{BASE_URL}/api/orders", headers=admin_h, json=payload, timeout=15)
        assert r.status_code in (200, 201), r.text
        order = r.json()
        oid = order["id"]
        # manual verify (admin)
        v = requests.post(f"{BASE_URL}/api/orders/{oid}/verify", headers=admin_h, timeout=15)
        assert v.status_code == 200, v.text
        # fetch merchant balance; should be credited by merchant_receive
        rs2 = requests.get(f"{BASE_URL}/api/accounts/{m['id']}", headers=admin_h, timeout=10)
        assert rs2.status_code == 200
        after = float(rs2.json().get("balance") or 0)
        assert after > before, f"merchant balance not credited: before={before} after={after}"


# ---------- Profit snapshot excel ----------
class TestProfitSnapshotExcel:
    DATE = "2026-10-10"

    def test_admin_xlsx(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/reports/profit-snapshot?date={self.DATE}", headers=admin_h, timeout=30)
        assert r.status_code == 200, r.text
        assert XLSX_MIME in r.headers.get("content-type", "")
        assert r.content[:2] == b"PK", "xlsx body should start with PK zip signature"
        # verify sheets
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert "日度汇总" in wb.sheetnames
        assert "按商户明细" in wb.sheetnames

    def test_operator_forbidden(self, op_h):
        r = requests.get(f"{BASE_URL}/api/reports/profit-snapshot?date={self.DATE}", headers=op_h, timeout=15)
        assert r.status_code == 403, f"operator should be 403, got {r.status_code}"


# ---------- Orders excel export ----------
class TestOrdersExcel:
    def test_filter_status_paid(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/reports/orders?status=paid", headers=admin_h, timeout=30)
        assert r.status_code == 200, r.text
        assert XLSX_MIME in r.headers.get("content-type", "")
        assert r.content[:2] == b"PK"
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # ensure non-header rows all have status == 已支付 (col 7)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if row and row[6] is not None:
                assert row[6] == "已支付", f"non-paid row leaked: {row}"

    def test_filter_network(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/reports/orders?network=TRC20", headers=admin_h, timeout=30)
        assert r.status_code == 200
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        rows = list(wb.active.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if row and row[3] is not None:
                assert row[3] == "TRC20"

    def test_by_ids(self, admin_h):
        # grab 2 orders
        lr = requests.get(f"{BASE_URL}/api/orders?limit=2", headers=admin_h, timeout=15)
        assert lr.status_code == 200
        data = lr.json()
        items = data if isinstance(data, list) else (data.get("items") or [])
        if len(items) < 2:
            pytest.skip("need 2 orders")
        ids = f"{items[0]['id']},{items[1]['id']}"
        r = requests.get(f"{BASE_URL}/api/reports/orders?ids={ids}", headers=admin_h, timeout=30)
        assert r.status_code == 200
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        rows = [row for row in wb.active.iter_rows(min_row=2, values_only=True) if row and row[0]]
        assert len(rows) == 2, f"expected 2 rows, got {len(rows)}"

    def test_invalid_ids_400(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/reports/orders?ids=not-an-objectid", headers=admin_h, timeout=15)
        assert r.status_code == 400

    def test_operator_without_perm_403(self, op_h):
        r = requests.get(f"{BASE_URL}/api/reports/orders?status=paid", headers=op_h, timeout=15)
        assert r.status_code == 403


# ---------- Telegram daily report ----------
class TestDailyReport:
    def test_preview(self, admin_h):
        r = requests.post(f"{BASE_URL}/api/notify/daily-report?date=2026-10-10", headers=admin_h, timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert set(["text", "telegram_configured", "sent", "date"]).issubset(d.keys())
        assert d["telegram_configured"] is False
        assert d["sent"] is False
        assert d["date"] == "2026-10-10"
        for kw in ["订单数", "收款总额", "平台利润"]:
            assert kw in d["text"], f"missing '{kw}' in text: {d['text']}"

    def test_future_date_400(self, admin_h):
        r = requests.post(f"{BASE_URL}/api/notify/daily-report?date=2099-01-01", headers=admin_h, timeout=15)
        assert r.status_code == 400

    def test_operator_forbidden(self, op_h):
        r = requests.post(f"{BASE_URL}/api/notify/daily-report?date=2026-10-10", headers=op_h, timeout=15)
        assert r.status_code == 403
